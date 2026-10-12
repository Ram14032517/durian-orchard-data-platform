"""Phase 11C offline audit. Imports existing evidence; never fetches or infers targets."""
import copy
import csv
import io
import json
import math
from pathlib import Path
from openpyxl import load_workbook
from tools import build_tha_mai_spatial_pilot as p
from tools.extract_chanthaburi_district_yield import FIELDS, UNITS, yield_check, difference

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'research_data/internal_research/chanthaburi_district_yield'
OLD = BASE / 'phase11b_20261012_v2'
OUT = BASE / 'phase11c_20261012_v1'
CSV_SHA = 'f24a9f0d62c01086eb8526316eae161a69ad625cacf8420238b35f60469dda0c'
OAE = ROOT / 'research_data/external/oae_2026-09-17/durian_province.xlsx'
OAE_SHA = '6a3ae44026b5ef8037a8bb091852281d890168a8d795c9c44d5cde9a48613d4c'


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def validate(rows):
    keys = [(r['district_code'], r['year_be'], r['source_version']) for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError('Duplicate district-year-source version')
    crosswalk = {r['district_code']: r['table_district_name'] for r in read(OLD / 'district_crosswalk.json')}
    for r in rows:
        if crosswalk.get(r['district_code']) != r['district_name']:
            raise ValueError('Crosswalk mismatch')
        if r['province_code'] != '22' or r['commodity'] != 'ทุเรียน':
            raise ValueError('Geographic/crop scope')
        if r['year_ce'] != r['year_be'] - 543 or not r['year_basis']:
            raise ValueError('Year definition missing/inconsistent')
        if r['year_basis'].startswith('unverified') and r['acceptance_scope'] != 'source_transcription_only':
            raise ValueError('Unverified year cannot certify actuals')
        if r['finality'] == 'final' and not r.get('status_evidence'):
            raise ValueError('Final without evidence')
        if r['source_version'] != r['source_sha256'] or not r['source_anchor']:
            raise ValueError('Source lineage missing')
        for f in FIELDS:
            v = r[f]
            if r['units'].get(f) != UNITS[f]:
                raise ValueError('Unit mismatch')
            if v is not None and (not math.isfinite(v) or v < 0):
                raise ValueError('Negative/nonfinite value')
            if r['missingness'].get(f) != ('present' if v is not None else 'missing_in_source'):
                raise ValueError('Missingness mismatch')
        if yield_check(r)['status'] == 'inconsistent':
            raise ValueError('Arithmetic inconsistency')
    return True


def target_complete(r):
    return all(r[f] is not None for f in FIELDS[1:])


def candidate_rows():
    if p.sha(OUT / 'candidate_standing_area.csv') != CSV_SHA:
        raise ValueError('CSV hash mismatch')
    source = read(OUT / 'candidate_standing_source.json')
    names = {r['table_district_name']: r['district_code'] for r in read(OLD / 'district_crosswalk.json')}
    reader = csv.DictReader(io.StringIO((OUT / 'candidate_standing_area.csv').read_text(encoding='utf-8-sig')))
    if reader.fieldnames != ['ปี', 'อำเภอ', 'เนื้อที่ยืนต้นทุเรียน', 'หน่วย']:
        raise ValueError('Unexpected CSV schema')
    rows = []
    for ordinal, raw in enumerate(reader, 1):
        if raw['หน่วย'].strip() != 'ไร่' or raw['อำเภอ'].strip() not in names:
            raise ValueError('Unverified unit/name')
        name = raw['อำเภอ'].strip()
        year = int(raw['ปี'])
        rows.append(dict(province_code='22', province_name='จันทบุรี', district_code=names[name],
            district_name=name, year_be=year, year_ce=year-543, commodity='ทุเรียน',
            year_basis='unverified_year_label_only', standing_area_rai=float(raw['เนื้อที่ยืนต้นทุเรียน'].replace(',', '')),
            bearing_area_rai=None, production_tonnes=None, reported_yield_kg_rai=None,
            units=UNITS.copy(), missingness={f: 'present' if f == FIELDS[0] else 'missing_in_source' for f in FIELDS},
            source_version=CSV_SHA, source_sha256=CSV_SHA, source_url=source['source_url'],
            source_agency='สำนักงานเกษตรจังหวัดจันทบุรี (catalog owner/source: ทะเบียนเกษตรกร)',
            document_title=source['catalog_title'], publication_date=None, revision_date=None,
            finality='unverified', acceptance_scope='source_transcription_only', training_eligible=False,
            source_anchor={'csv_record_1based_excluding_header_and_blank_lines': ordinal,
                           'csv_line_end_1based': reader.line_num, 'columns': reader.fieldnames},
            quality_flags=['standing_only_not_yield', 'year_basis_unverified', 'finality_unverified',
                           'publication_date_unverified', 'registration_population_not_proven_equal_to_oae',
                           'license_scope_pending_public_hold'] +
                          (['historical_row_alignment_or_definition_review_needed'] if year in (2557,2558) else [])))
    validate(rows)
    return rows


def export():
    # Exclusive creates make a completed run immutable; no retry may replace its outputs.
    if (OUT / 'verified_district_year.json').exists():
        raise FileExistsError('Refuse overwrite existing audit')
    p.previous.verify_seal(ROOT, str(OLD.relative_to(ROOT)))
    before = read(OUT / 'preflight_protected_hashes.json')
    if any(not (ROOT/k).is_file() or p.sha(ROOT/k) != v for k,v in before.items()):
        raise ValueError('Protected input changed before export')
    if p.sha(OAE) != OAE_SHA:
        raise ValueError('OAE hash mismatch')
    manifest = read(OLD / 'source_manifest.json')
    verified = copy.deepcopy(read(OLD / 'district_year.json')['rows'])
    for r in verified:
        r.update(source_version=manifest['sha256'], source_sha256=manifest['sha256'],
                 source_url=manifest['approved_url'], source_agency=manifest['table_attribution'],
                 document_title=manifest['document_title'], commodity='ทุเรียน', revision_date=None,
                 provenance='carried_forward_phase11b_not_new_discovery')
    validate(verified)
    candidates = candidate_rows()
    assert len(candidates) == 110 and len(verified) == 20
    overlap = [{ 'district_code': r['district_code'], 'year_be': r['year_be'],
        'field': 'standing_area_rai', 'pdf_value': r['standing_area_rai'],
        'catalog_value': next(c['standing_area_rai'] for c in candidates
            if (c['district_code'],c['year_be']) == (r['district_code'],r['year_be']))} for r in verified]
    for r in overlap:
        r['conflict'] = r['pdf_value'] != r['catalog_value']
    coverage = []
    for year in range(2557,2568):
        for code in map(str,range(2201,2211)):
            coverage.append({'year_be':year,'year_ce':year-543,'district_code':code,
                'reported_target_transcription_available': any(r['year_be']==year and r['district_code']==code and target_complete(r) for r in verified),
                'standing_only_candidate_available':any(r['year_be']==year and r['district_code']==code for r in candidates),
                'final_actual_verified':False, 'year_definition_verified':False, 'publication_vintage_verified':False})
    reconciliation = copy.deepcopy(read(OLD / 'reconciliation.json'))
    w = load_workbook(OAE,read_only=True,data_only=True)
    references = {}
    for row_num, cells in enumerate(w['durian'].iter_rows(min_row=2,values_only=True),2):
        year,crop,province,commodity,subcom,item,value,unit,data_date = cells
        if province=='จันทบุรี' and commodity=='ทุเรียน' and item=='เนื้อที่ยืนต้น' and 2557<=int(year)<=2567:
            if unit!='ไร่' or int(year) in references:
                raise ValueError('OAE unit/duplicate reference')
            references[int(year)] = {'value':value,'sheet':'durian','cell':f'G{row_num}', 'data_date':str(data_date)}
    w.close()
    standing_reconciliation = []
    for year in sorted({r['year_be'] for r in candidates}):
        total=sum(r['standing_area_rai'] for r in candidates if r['year_be']==year)
        ref=references.get(year)
        standing_reconciliation.append({'year_be':year,'catalog_district_sum_rai':total,
            'oae_reference':ref,'difference':difference(total,ref['value']) if ref else None,
            'interpretation':'different_source_population/vintage_possible; no scaling; data_date is not publication date'})
    lineage = []
    for r in verified+candidates:
        for field in ['province_code','district_code','district_name','year_be','year_ce','year_basis', *FIELDS,
                      'commodity','finality','publication_date','revision_date']:
            lineage.append({'district_code':r['district_code'],'year_be':r['year_be'],'source_version':r['source_version'],
                'field':field,'value':r[field],'source_sha256':r['source_sha256'],'source_url':r['source_url'],
                'anchor':r['source_anchor'], 'method':
                'not_reported_or_not_verified_null' if r[field] is None else
                'year_be_minus_543_not_year_definition_proof' if field=='year_ce' else
                'phase11b_sealed_admin_crosswalk' if field in ('province_code','district_code') else
                'explicit_unverified_status_no_inference' if field in ('finality','year_basis') else
                'inherited_visual_transcription_phase11b' if r in verified else 'CSV_named_column_exact_transcription'})
    packages=read(OUT / 'catalog_discovery.json')['response']['result']['results']
    inventory=[]
    for x in packages:
        if x['name'] in ('dataset_10_05','dataset_10_06','dataset_10_07'):
            inventory.append({'source_id':x['name'],'url':f"https://chanthaburi.gdcatalog.go.th/dataset/{x['name']}",
                'title':x['title'],'source':x.get('data_source'),'grain':x.get('geo_coverage'),
                'evidence':'captured CKAN package metadata','resources':x['resources'],
                'decision':'downloaded_standing_only_exclude_yield' if x['name']=='dataset_10_07' else 'province_only_not_downloaded',
                'publication_date':None,'metadata_created':x.get('metadata_created'),'metadata_modified':x.get('metadata_modified')})
    inventory += read(OUT / 'official_web_checks.json')
    qc={'verified_transcription_rows':20,'verified_yield_years_be':[2566,2567], 'complete_10_district_target_years':2,
        'new_yield_years':0,'verified_final_actual_years':0,'standing_only_candidate_rows':110,
        'standing_only_years_be':list(range(2557,2568)),'missing_target_years_be':list(range(2557,2566)),
        'ratio_checks':[{'district_code':r['district_code'],'year_be':r['year_be'],**yield_check(r)} for r in verified],
        'overlap_conflicts':sum(r['conflict'] for r in overlap),'historical_row_alignment_warning_years_be':[2557,2558],
        'warning_evidence':'ท่าใหม่ standing 1667/1583 vs แหลมสิงห์ 66224/66653; abrupt placement change in 2559. Cause unverified; no swaps/corrections.',
        'license_status':'internal_only_public_hold; catalog license label is not historical publication evidence',
        'acceptance':'source_transcription_only_not_final_actual_or_historical_PIT',
        'search_scope':'bounded official search, not exhaustive; no usable additional district targets verified',
        'network_downloads':'one standing CSV + CKAN metadata; no new PDF downloaded; Phase11B PDF reused read-only'}
    products={'verified_district_year.json':{'rows':verified,'public_release':'HOLD','scope':'verified transcription, NOT verified actuals'},
              'candidate_unverified_records.json':{'rows':candidates,'target_eligible':False},
              'coverage_matrix.json':coverage,'oae_reconciliation.json':{'phase11b_target_reference':reconciliation,
                    'phase11b_source_path':str(OLD.relative_to(ROOT)),'standing_candidates':standing_reconciliation},
              'cross_source_comparison.json':overlap,'source_discovery_inventory.json':inventory,
              'field_lineage.json':lineage,'qc_report.json':qc,
              'data_dictionary.json':{'row_key':['district_code','year_be','source_version'],'numeric_units':UNITS,
                 'source_version':'SHA256 of exact resource bytes; not asserted official edition number',
                 'missing_values':'null means not reported/unverified, never zero',
                 'reported_yield_kg_rai':'source value only; computed ratio retained separately in QC',
                 'year_basis':'unverified_year_label_only, no calendar/harvest assertion',
                 'finality':'unverified unless original status evidence exists',
                 'publication_date':'null; retrieval/catalog timestamps must not substitute',
                 'revision_date':'null; metadata_modified is resource metadata only',
                 'csv_anchor':'record ordinal excludes header and blank lines; line end is DictReader.line_num',
                 'source_agency':'source attribution, not proof of finality',
                 'quality_flags':'static footprint not joined; public HOLD, no ML eligibility'},
              'source_manifest.json':{'phase11b_source':manifest,'candidate_source':read(OUT/'candidate_standing_source.json'),
                 'inputs':{str(path.relative_to(ROOT)):p.sha(path) for path in
                   [OLD/'district_year.json',OLD/'output_manifest.json',OLD/'field_lineage.json',OLD/'district_crosswalk.json',
                    OLD/'source.pdf',OAE, OUT/'candidate_standing_area.csv',OUT/'catalog_discovery.json',OUT/'official_web_checks.json']},
                 'transformation_code_sha256':p.sha(Path(__file__)),'public_release':'HOLD'}}
    for name,value in products.items():
        p.write(OUT/name,value)
    after={k:p.sha(ROOT/k) if (ROOT/k).is_file() else None for k in before}
    mismatches=[k for k,v in before.items() if after[k]!=v]
    if mismatches:
        raise ValueError(f'Protected mismatch: {mismatches}')
    p.write(OUT/'protected_verification.json',{'checked_files':len(before),'mismatches':mismatches,
             'production_sha256':p.sha(ROOT/p.previous.MODEL),'production_expected':p.previous.MODEL_HASH,
             'unchanged':True,'public_release':'HOLD'})
    print(json.dumps({'rows':len(verified),'candidate_rows':len(candidates),'new_yield_years':0,
                     'overlap_conflicts':qc['overlap_conflicts'],'protected_files':len(before)},ensure_ascii=False))


if __name__ == '__main__':
    export()
