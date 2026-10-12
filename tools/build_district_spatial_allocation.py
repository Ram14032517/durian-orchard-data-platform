"""Retrospective allocation, not measured district yield or a forecasting model."""
import argparse
import copy
import csv
import json
import math
from pathlib import Path
from statistics import mean
from tools.build_chanthaburi_advisor_eda import ROOT,P9,P10,P11,MODEL,sha,read,rows,seal_check,table,e

P11C=ROOT/'research_data/internal_research/chanthaburi_district_yield/phase11c_20261012_v1'
FINAL=ROOT/'research_data/internal_research/final_advisor/final_20261012_v1'
CODES=[str(c) for c in range(2201,2211)]
YEARS=list(range(2559,2566))
CONFIG=dict(calibration_years_be=[2566,2567],allocation_years_be=YEARS,
    fixed_share='Equal-year average of district production / sum of ten district productions',
    area_adjusted='r_d=mean(P_d,t/A_d,t); w_d,y=A_d,y*r_d / sum_j(A_j,y*r_j)',
    validation_pairs=[[2566,2567],[2567,2566]],
    validation_target='district reported production tonnes, not kg/rai yield',
    area_gate='Fail-closed on unresolved calibration standing-area source conflict; no source substitution/drop',
    provincial_total='Known OAE total for allocation year; conditional allocation, not independent forecasting',
    gis='context only; never used as historical area or allocation weights',
    status='Internal hypothetical retrospective allocation; public HOLD; prohibited ML input')

def save(p,v):
    with p.open('x',encoding='utf-8') as f:json.dump(v,f,ensure_ascii=False,indent=2,allow_nan=False)
def write(p,v):
    with p.open('x',encoding='utf-8') as f:f.write(v)
def csvsave(p,data):
    with p.open('x',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)

def fixed_weights(reported,calibration,target):
    assert target not in calibration,'Target year cannot enter calibration'
    shares=[]
    for year in calibration:
        rr=[r for r in reported if r['year_be']==year]
        assert len(rr)==10 and {r['district_code'] for r in rr}==set(CODES)
        assert all(r['production_tonnes'] is not None and math.isfinite(r['production_tonnes']) and r['production_tonnes']>=0 for r in rr)
        total=sum(r['production_tonnes'] for r in rr);assert total>0
        shares.append({r['district_code']:r['production_tonnes']/total for r in rr})
    w={c:mean(s[c] for s in shares) for c in CODES}
    assert abs(sum(w.values())-1)<1e-12
    return w

def area_gate(conflicts,calibration,target):
    return not any(r['year_be'] in set(calibration+[target]) for r in conflicts)

def tests(reported,conflicts):
    w=fixed_weights(reported,[2566],2567)
    mutant=copy.deepcopy(reported)
    for r in mutant:
        if r['year_be']==2567:r['production_tonnes']+=1234567
    assert w==fixed_weights(mutant,[2566],2567)
    try:fixed_weights(reported,[2566,2567],2567)
    except AssertionError:pass
    else:raise AssertionError('target calibration accepted')
    try:fixed_weights([r for r in reported if not (r['year_be']==2566 and r['district_code']=='2202')],[2566],2567)
    except AssertionError:pass
    else:raise AssertionError('missing district accepted')
    assert not area_gate(conflicts,[2566],2567) and not area_gate(conflicts,[2567],2566)
    assert abs(sum(123456*v for v in w.values())-123456)<1e-8
    return ['target production mutation leaves weights unchanged','target calibration rejected','missing district rejected','area conflicts fail-closed both directions','weights/allocated total conservation']

def build(out):
    assert not out.exists(),'Refuse overwrite'
    protected=read(ROOT/'research_data/internal_research/chanthaburi_district_yield/phase11e_20261012_v1/preflight_protected_hashes.json')
    assert all(sha(ROOT/k)==v for k,v in protected.items()),'Protected mismatch'
    for folder in (P9,P10,P11,P11C,FINAL):seal_check(folder)
    production=sha(MODEL);assert production=='b4b9e0985e6ba5e6ee371d8988cad5e65dee031b345eb0cf13a8695f91658680'
    from pypdf import PdfReader
    prior=PdfReader(FINAL/'final_advisor_report_th.pdf')
    assert len(prior.pages)==10 and 'Scenario' in prior.pages[6].extract_text()
    reported=read(P11/'district_year.json')['rows']
    assert len(reported)==20 and {(r['district_code'],r['year_be']) for r in reported}=={(c,y) for c in CODES for y in (2566,2567)}
    names={r['district_name']:r['district_code'] for r in reported};assert len(names)==10
    areas={};area_anchors={}
    for rownum,r in enumerate(rows(P11C/'candidate_standing_area.csv'),2):
        name=r['อำเภอ'];year=int(r['ปี']);value=r['เนื้อที่ยืนต้นทุเรียน'].replace(',','')
        assert name in names and r['หน่วย']=='ไร่' and value!=''
        key=(names[name],year);assert key not in areas
        areas[key]=float(value);area_anchors[key]=f'CSV row {rownum}, เนื้อที่ยืนต้นทุเรียน'
        assert math.isfinite(areas[key]) and areas[key]>0
    assert len(areas)==110 and set(areas)=={(c,y) for c in CODES for y in range(2557,2568)}
    conflicts=[]
    for r in reported:
        v=areas[(r['district_code'],r['year_be'])]
        if v!=r['standing_area_rai']:
            conflicts.append(dict(district_code=r['district_code'],district_name=r['district_name'],year_be=r['year_be'],phase11b_standing_rai=r['standing_area_rai'],phase11c_standing_rai=v,difference_rai=v-r['standing_area_rai'],status='unresolved_source_conflict',phase11b_anchor='PDF page 5 Table 2 row '+str(r['source_anchor']['district_row_ordinal_excluding_province']),phase11c_anchor=area_anchors[(r['district_code'],r['year_be'])]))
    assert len(conflicts)==read(P11C/'qc_report.json')['overlap_conflicts']==2
    province={int(r['year_ce'])+543:float(r['production_tonnes']) for r in rows(P9/'province_year.csv') if r['province_code']=='22'}
    from openpyxl import load_workbook
    wbpath=ROOT/'research_data/external/oae_2026-09-17/durian_province.xlsx'
    workbook=load_workbook(wbpath,read_only=True,data_only=True);anchors={}
    for index,r in enumerate(workbook['durian'].iter_rows(min_row=2,values_only=True),2):
        year,crop,prov,commodity,subcom,item,value,unit,date=r
        if prov=='จันทบุรี' and commodity=='ทุเรียน' and item=='ผลผลิต' and int(year) in YEARS+[2566,2567]:
            assert unit=='ตัน' and float(value)==province[int(year)]
            anchors[int(year)]=dict(sheet='durian',cell=f'G{index}',data_date=str(date))
    workbook.close();assert len(anchors)==9
    reconciliation=[]
    for year in (2566,2567):
        total=sum(r['production_tonnes'] for r in reported if r['year_be']==year)
        reconciliation.append(dict(year_be=year,district_reported_sum_tonnes=total,oae_province_total_tonnes=province[year],difference_tonnes=total-province[year]))
        assert total==province[year]
    checks=tests(reported,conflicts)+['110 area keys/units/positive values','20 reported district-year keys','9 OAE production cells/units','reported district totals equal OAE in two years','prior dataset seals and PDF verified']
    inputs=[P11/'district_year.json',P11/'field_lineage.json',P11/'source_manifest.json',P11/'output_manifest.json',P11C/'candidate_standing_area.csv',P11C/'candidate_standing_source.json',P11C/'cross_source_comparison.json',P11C/'qc_report.json',P11C/'output_manifest.json',P9/'province_year.csv',P9/'source_manifest.json',P9/'output_manifest.json',P10/'qc_summary.json',P10/'release_metadata.json',P10/'output_manifest.json',FINAL/'final_advisor_report_th.pdf',FINAL/'output_manifest.json',wbpath,Path(__file__),ROOT/'tools/build_chanthaburi_advisor_eda.py']
    before={str(p.relative_to(ROOT)):sha(p) for p in inputs}
    out.mkdir(parents=True);save(out/'config.json',CONFIG)
    estimates=[];params=[];weight=fixed_weights(reported,[2566,2567],2559)
    for c in CODES:
        params.append(dict(district_code=c,district_name=next(r['district_name'] for r in reported if r['district_code']==c),fixed_share=weight[c],share_2566=fixed_weights(reported,[2566],2567)[c],share_2567=fixed_weights(reported,[2567],2566)[c],area_adjusted_rate_tonnes_per_standing_rai=None,area_adjusted_status='blocked_unresolved_calibration_area_conflict'))
    for year in YEARS:
        for method in ('fixed_share','area_adjusted_share'):
            assert year not in [2566,2567]
            for c in CODES:
                blocked=method=='area_adjusted_share';assert not area_gate(conflicts,[2566,2567],year) if blocked else True
                estimates.append(dict(province_code='22',district_code=c,district_name=next(r['district_name'] for r in reported if r['district_code']==c),year_be=year,year_ce=year-543,method=method,estimated_production_tonnes=None if blocked else province[year]*weight[c],allocated_share=None if blocked else weight[c],oae_province_total_tonnes=province[year],standing_area_rai=None if not blocked else areas[(c,year)],standing_area_role='not_used' if not blocked else 'candidate_not_used_gate_failed',calibration_years_be='2566;2567',record_type='blocked_estimate' if blocked else 'hypothetical_retrospective_allocation',quality_status='blocked_unresolved_area_vintage' if blocked else 'arithmetic_pass_not_historical_accuracy_verified',use_for_ml=False,public_release='HOLD',oae_anchor='durian!'+anchors[year]['cell'],calibration_source_sha256=sha(P11/'district_year.json')))
        assert abs(sum(r['estimated_production_tonnes'] for r in estimates if r['year_be']==year and r['method']=='fixed_share')-province[year])<1e-8
    validations=[];summaries=[]
    for cal,target in ((2566,2567),(2567,2566)):
        w=fixed_weights(reported,[cal],target)
        for method in ('fixed_share','area_adjusted_share'):
            blocked=method=='area_adjusted_share'; rr=[]
            for c in CODES:
                actual=next(r for r in reported if r['year_be']==target and r['district_code']==c)
                pred=None if blocked else province[target]*w[c]
                r=dict(calibration_year_be=cal,target_year_be=target,district_code=c,district_name=actual['district_name'],method=method,reported_production_tonnes=actual['production_tonnes'],estimated_production_tonnes=pred,error_tonnes=None if blocked else pred-actual['production_tonnes'],absolute_error_tonnes=None if blocked else abs(pred-actual['production_tonnes']),oae_target_total_tonnes=province[target],weights_include_target_production=False,temporal_direction='forward' if cal<target else 'reverse_retrospective_not_forecast',status='blocked_unresolved_area_conflict' if blocked else 'conditional_allocation_test_reported_reference',reported_finality='unverified')
                rr.append(r);validations.append(r)
            if blocked: metric=dict(mae_tonnes=None,rmse_tonnes=None,bias_tonnes=None,wape_pct=None,max_absolute_error_tonnes=None,worst_district=None)
            else:
                errors=[r['error_tonnes'] for r in rr];total=sum(r['reported_production_tonnes'] for r in rr)
                worst=max(rr,key=lambda r:r['absolute_error_tonnes'])
                metric=dict(mae_tonnes=mean(abs(v) for v in errors),rmse_tonnes=math.sqrt(mean(v*v for v in errors)),bias_tonnes=mean(errors),wape_pct=100*sum(abs(v) for v in errors)/total,max_absolute_error_tonnes=worst['absolute_error_tonnes'],worst_district=worst['district_name'])
                assert abs(sum(errors))<1e-8
            summaries.append(dict(calibration_year_be=cal,target_year_be=target,method=method,n=0 if blocked else 10,status='blocked' if blocked else 'computed',**metric))
    assert len(estimates)==140 and sum(r['estimated_production_tonnes'] is not None for r in estimates)==70
    assert len(validations)==40 and all(r['estimated_production_tonnes'] is None for r in validations if r['method']=='area_adjusted_share')
    assert len({(r['district_code'],r['year_be'],r['method']) for r in estimates})==140
    checks+=['70 allocated rows/70 blocked rows remain distinct','7 province totals conserved','both validation directions target excluded','no missing estimates filled','no estimated yield or GIS weights']
    csvsave(out/'estimated_district_production_2559_2565.csv',estimates)
    csvsave(out/'reported_district_data_2566_2567.csv',[dict(province_code='22',district_code=r['district_code'],district_name=r['district_name'],year_be=r['year_be'],year_ce=r['year_ce'],reported_production_tonnes=r['production_tonnes'],reported_bearing_area_rai=r['bearing_area_rai'],reported_standing_area_rai=r['standing_area_rai'],reported_yield_kg_rai=r['reported_yield_kg_rai'],record_type='reported_source_transcription',finality='unverified',source_sha256=r['source_anchor']['pdf_sha256'],source_page=5,source_table=2,source_row=r['source_anchor']['district_row_ordinal_excluding_province']) for r in reported])
    csvsave(out/'validation_details.csv',validations);csvsave(out/'validation_summary.csv',summaries);csvsave(out/'weight_parameters.csv',params);csvsave(out/'area_qc_conflicts.csv',conflicts)
    save(out/'source_manifest.json',dict(input_hashes=before,oae_cell_anchors=anchors,area_source=read(P11C/'candidate_standing_source.json'),district_source_anchors=[r['source_anchor'] for r in reported],oae_reconciliation=reconciliation,public_release='HOLD',publication_vintage='unverified'))
    f=lambda v:f'{v:,.2f}'
    pages=[]
    def page(title,body):return '<section class="page"><header>Selected Topics in AI <span>จันทบุรี / Internal Only</span></header><h1>'+title+'</h1>'+body+f'<footer>ค่าจัดสรรสมมติ ไม่ใช่ข้อมูลผลผลิตรายอำเภอที่สำรวจจริง <span>{len(pages)+1} / 4</span></footer></section>'
    pages.append(page('ทดลองจัดสรรผลผลิตจังหวัดลงอำเภอ','<p class="lead">ต่อยอดรายงานอาจารย์ด้วยการจัดสรรเชิงสมมติ ไม่ทำ EDA หรือฝึก ML ซ้ำ</p><p>ประมาณปี พ.ศ.2559–2565 (2016–2022) เท่านั้น หน่วยผลผลิตรวม “ตัน” ไม่คำนวณ Yield กก./ไร่ เพราะไม่มีพื้นที่ให้ผลรายอำเภอครบย้อนหลัง</p><h2>1. Fixed Share</h2><p>s<sub>d,t</sub> = P<sub>d,t</sub> / Σ<sub>j</sub>P<sub>j,t</sub> ในแต่ละปีอ้างอิง 2566 และ 2567<br>w<sub>d</sub> = (s<sub>d,2566</sub> + s<sub>d,2567</sub>) / 2<br>P̂<sub>d,y</sub> = P<sub>จังหวัด,y</sub> × w<sub>d</sub></p><p>ใช้ค่าเฉลี่ยสัดส่วนให้น้ำหนักสองปีเท่ากัน ไม่ใช่สัดส่วนจากผลผลิตรวมสองปี ผลรวมรายอำเภอตรงยอดจังหวัดตามสูตรโดยอัตโนมัติ แต่ไม่ได้พิสูจน์ว่าการแบ่งแต่ละอำเภอถูกต้อง</p><h2>2. Area-adjusted Share</h2><p>แนวทางที่ต้องการทดสอบ: r<sub>d</sub> = mean(P<sub>d,t</sub>/A<sub>d,t</sub>) จากปีอ้างอิง<br>w<sub>d,y</sub> = A<sub>d,y</sub>r<sub>d</sub> / Σ<sub>j</sub>A<sub>j,y</sub>r<sub>j</sub><br>A คือพื้นที่ยืนต้น ไม่ใช่พื้นที่ให้ผล อัตรา P/A จึงไม่ใช่ Yield ที่ OAE รายงาน</p><div class="notice"><strong>Area-adjusted ไม่ผ่าน QC ต้นทาง</strong><p>พื้นที่ยืนต้นขลุงต่างกันระหว่าง Phase11B/11C ทั้งสองปีอ้างอิง ไม่ทราบรุ่น/นิยามที่ทำให้ต่าง จึงไม่คำนวณอัตราหรือน้ำหนัก ไม่เลือกต้นทางหนึ่งแทนอีกต้นทางเอง ไม่ Drop ขลุงเพื่อปรับ Normalize ใหม่ เก็บ 70 ค่าประมาณเป็น null พร้อมเหตุผล</p></div>'))
    metricrows=[[r['calibration_year_be'],r['target_year_be'],'Fixed' if r['method']=='fixed_share' else 'Area-adjusted',f(r['mae_tonnes']) if r['mae_tonnes'] is not None else 'ไม่คำนวณ',f(r['rmse_tonnes']) if r['rmse_tonnes'] is not None else 'QC blocked',f(r['wape_pct']) if r['wape_pct'] is not None else '-'] for r in summaries]
    fixed=[r for r in summaries if r['method']=='fixed_share']
    detail=[]
    for c in CODES:
        rr=[r for r in validations if r['district_code']==c and r['method']=='fixed_share']
        detail.append([rr[0]['district_name'],f(rr[0]['error_tonnes']),f(rr[1]['error_tonnes'])])
    pages.append(page('ตรวจข้ามปีโดยไม่ใช้ Target สร้างน้ำหนัก',table(['ปีอ้างอิง','ปีตรวจ','วิธี','MAE ตัน','RMSE ตัน','WAPE %'],metricrows)+'<p>MAE/RMSE เฉลี่ยรายอำเภอ n=10 ต่อทิศทาง WAPE = Σ|Error| / Σผลผลิตที่รายงาน ×100 Error = จัดสรร − รายงาน</p>'+table(['อำเภอ','Error: 2566 → 2567','Error: 2567 → 2566'],detail)+f'<p>กรณีผิดมากสุด: 2566→2567 {fixed[0]["worst_district"]} {f(fixed[0]["max_absolute_error_tonnes"])} ตัน; สลับปี {fixed[1]["worst_district"]} {f(fixed[1]["max_absolute_error_tonnes"])} ตัน</p><p>ใช้ยอดรวมจังหวัดปีเป้าหมายที่รู้แล้วเป็นเงื่อนไข ไม่ใช้ผลผลิตรายอำเภอปีเป้าหมายในการสร้างสัดส่วน จึงไม่ใช่การพยากรณ์ผลผลิตจังหวัดอิสระ Bias รวมใกล้ศูนย์เกิดจากบังคับยอดรวม ไม่ใช่หลักฐานความแม่น</p><div class="notice">2567→2566 เป็นการถ่ายโอนย้อนเวลา ไม่ใช่ Forecast ที่ทำได้ในอดีต มีเพียงสองปีที่พึ่งพากัน ไม่ใช้พิสูจน์ความแม่นของปี 2559–2565 ไม่ใช้ผลทดสอบเลือกวิธีที่ชนะ</div>'))
    hist=[]
    for c in CODES:
        hist.append([next(r['district_name'] for r in reported if r['district_code']==c),f'{weight[c]*100:.2f}']+[f'{province[y]*weight[c]:,.0f}' for y in YEARS])
    hist.append(['ยอดจังหวัด','100.00']+[f'{province[y]:,.0f}' for y in YEARS])
    pages.append(page('ค่าประมาณย้อนหลัง: Fixed Share เท่านั้น','<p><strong>ตารางนี้ไม่ใช่ข้อมูลผลผลิตรายอำเภอที่สำรวจจริง</strong> มี 70 ค่า จากยอดจังหวัด OAE × ค่าเฉลี่ยสัดส่วนปี 2566–2567 ทุกค่าติดป้าย hypothetical_retrospective_allocation และ use_for_ml=False</p><div class="wide">'+table(['อำเภอ','Share %']+[str(y) for y in YEARS],hist)+'</div><p>หน่วยตัน ปัดจำนวนเต็มเฉพาะในภาพ ตาราง CSV เก็บความละเอียดคำนวณ ผลรวมหลังปัดในภาพอาจต่างยอดจังหวัดเล็กน้อย ไม่มีการปรับตัวเลขรายอำเภอให้ยอดที่ปัดตรงกัน</p><h2>สมมติฐานที่หนักที่สุด</h2><p>สัดส่วนผลผลิตของแต่ละอำเภอในปี 2566–2567 ถูกสมมติให้แทนอดีต 7 ปี ซึ่งอาจไม่จริงเมื่อพื้นที่ปลูก อายุสวน สัดส่วนพื้นที่ให้ผล การจัดการและฤดูเปลี่ยน การใช้ปีหลังย้อนจัดสรรเป็นข้อสมมติที่ตั้งใจ ไม่อ้าง Historical PIT หรือการพยากรณ์ล่วงหน้า</p><h2>ข้อมูลที่ไม่ถูกเติม</h2><p>ค่าผลผลิตจริงรายอำเภอปี 2559–2565 ยัง Missing ไม่มีการสร้างค่านั้นเป็น Actual Area-adjusted ยัง Missing เพราะ QC blocked ไม่คำนวณ Yield จากพื้นที่ยืนต้นแทนพื้นที่ให้ผล และไม่แจกข้อมูลนี้ให้สวนหรือ Polygon</p><h2>GIS ใช้เพื่อบริบทเท่านั้น</h2><p>Pure A403 จันทบุรีปี 2568 และดินท่าใหม่ที่ผ่าน Policy v2 แสดงความต่างของหน่วยพื้นที่ ไม่เข้าการคำนวณน้ำหนัก ไม่แทนพื้นที่ปลูกย้อนหลัง ไม่ Join กับ Yield และไม่ใช้สรุปว่าดิน/อุณหภูมิเป็นเหตุของผลผลิต</p>'))
    pages.append(page('QC แหล่งข้อมูล และข้อจำกัด',table(['ปี / อำเภอ','พื้นที่ Phase11B ไร่','พื้นที่ Phase11C ไร่','ต่าง ไร่'],[[str(r['year_be'])+' ขลุง',f(r['phase11b_standing_rai']),f(r['phase11c_standing_rai']),f(r['difference_rai'])] for r in conflicts])+'<p>Phase11C มาจากทะเบียนเกษตรกรในบัญชีข้อมูลจังหวัด มี Candidate Standing Area 110 แถว 2557–2567 ครบ 10 อำเภอ ทั้งค่าหน่วย/คีย์ตรวจได้ แต่ Vintage/นิยามยังไม่ยืนยัน และปี2557–2558มี Row-alignment warning เดิม ไม่อยู่ในช่วงจัดสรร ไม่มีการสลับหรือแก้ค่าดังกล่าว</p><h2>Source และการตรวจ</h2><p>[1] OAE Workbook เดิม Sheet durian: Production ตัน ตรวจ 9 เซลล์ปี2559–2567 เทียบ Phase9D<br>[2] Phase11B v2: PDF หน้า5 Table2 ครบ20แถว ใช้ค่ารายงาน ไม่รับรอง Final Actual นิยามปี/Publication Date ยังไม่ยืนยัน<br>[3] Phase11C v1: candidate_standing_area.csv และ Cross-source Comparison เก็บค่าขัดแย้งเดิม<br>[4] รายงานอาจารย์เดิมและ Phase10F QC ใช้อ้างหน่วยพื้นที่เท่านั้น</p><p>ผลรวม Production รายอำเภอปี2566–2567ตรง OAE แต่ผลรวมตรงไม่ยืนยันค่าทุกอำเภอ Tests เฉพาะงานตรวจ Target exclusion, Missing keys, Weights, Conservation และ Fail-closed Area Gate ไม่มีการฝึก ML/Overlay/ดาวน์โหลด หรือ Repo-wide Tests ที่ไม่เกี่ยวข้อง</p><h2>ไฟล์และการอ่านผล</h2><p>estimated_district_production_2559_2565.csv แยกจาก reported_district_data_2566_2567.csv; validation_details/summary เก็บตรวจข้ามปี; weight_parameters แสดงที่มาสัดส่วน; area_qc_conflicts เก็บปัญหา พิกัด/Geometry ไม่ถูก Export ช่องว่าง CSV คือ null ไม่ใช่ศูนย์</p><div class="notice"><p>ยังเปรียบเทียบว่า Area-adjusted ดีกว่า Fixed ไม่ได้เพราะ QC ไม่ผ่าน และยังไม่ยืนยันความแม่นย้อนหลังของ Fixed เก็บไว้ใช้สนทนาเรื่องสมมติฐานเท่านั้น ห้ามนำค่าประมาณไปฝึก ML คง Public HOLD</p></div>'))
    assert len(pages)==4
    css='''@page{size:A4;margin:0}*{box-sizing:border-box}body{margin:0;background:#e8eef0;color:#132d38;font:15px/1.65 "Leelawadee UI",Tahoma,sans-serif}.page{width:210mm;min-height:297mm;background:white;padding:15mm 17mm 18mm;margin:8mm auto;position:relative;break-after:page}.page:last-child{break-after:auto}header{font-size:12px;border-bottom:1px solid #b8cbd1;padding-bottom:8px}header span,footer span{float:right}h1{font-size:25px;line-height:1.35;margin:22px 0 18px}h2{font-size:18px;margin:17px 0 8px}p{margin:10px 0}.lead{font-size:18px}.notice{padding:10px 14px;background:#fff5e8;border-left:4px solid #b45309;margin:14px 0}table{border-collapse:collapse;width:100%;font-size:13px}td,th{padding:7px 6px;border-bottom:1px solid #dbe3e9;text-align:left}th{background:#edf3f5}.wide table{font-size:12px}.wide td,.wide th{padding:6px 3px;white-space:nowrap}.wide td:first-child{white-space:normal}footer{position:absolute;bottom:10mm;left:17mm;right:17mm;font-size:11px;color:#49616c;border-top:1px solid #dbe3e9;padding-top:6px}@media print{body{background:white}.page{margin:0;height:297mm;min-height:0}}'''
    write(out/'spatial_allocation_report_th.html','<!doctype html><html lang="th"><meta charset="utf-8"><title>การจัดสรรผลผลิตจันทบุรีแบบสมมติ</title><style>'+css+'</style>'+''.join(pages)+'</html>')
    write(out/'README_TH.md','''# ทดลอง Spatial Allocation จันทบุรี (Internal Only / Public HOLD)

อ่านรายงาน4หน้า ค่าประมาณไม่ใช่ Actual หรือ Yield กก./ไร่ ข้อมูลจริงที่รายงานแยก CSV ไม่เปลี่ยนต้นทาง

Fixed Share: mean(Pdistrict/Pdistrict_sum) จาก2566–2567 แล้วคูณยอด OAE จังหวัดปี2559–2565. Cross-year validation ใช้ปีอ้างอิงปีเดียว ไม่ใช้ Target ใน Calibration. ยอดจังหวัดปี Target รู้แล้ว จึงเป็น Conditional Allocation ไม่ใช่ Forecast. Reverse2567→2566ไม่ใช่ PIT.

Area-adjusted: ตั้งสูตร P/Standing Area แต่ไม่คำนวณ เนื่องจากขลุงพื้นที่สองฉบับขัดแย้งทั้งปีอ้างอิง ไม่เปลี่ยนแหล่ง ไม่ตัดอำเภอ ไม่ Normalize กลุ่มย่อย. CSV140แถวเป็นFixed70ค่ากับArea70null. No-fill. ข้อมูลพื้นที่2557–58มีWarningเดิมและไม่ใช้.

## CSV และหน่วย
- estimated_district_production_2559_2565.csv: grain District–Year–Method, tonnes, allocated_share fraction0–1, ปีBE/CE, calibrating years2566;2567, status/record_type/use_for_ml=False. standing_area_rai เฉพาะ Candidate ของวิธีBlocked ไม่ได้ใช้ในสูตร
- reported_district_data_2566_2567.csv: grain District–Year, reported Production tonnes/Bearing rai/Standing rai/Yield kg per bearing rai, PDF anchors/sourceSHA; Finality unverified
- validation_details.csv: grain Calibration–Target–Method–District; Error=Estimate−Report, all tonnes; Area method predictions/errors null
- validation_summary.csv: 4 rows; MAE/RMSE/Bias tonnes, WAPE percent, n=0 for blocked; no invented metrics
- weight_parameters.csv: 10districts, Fixedshares2566/67/mean fractions; area_rate null
- area_qc_conflicts.csv: original source values rai/anchors, no corrections

จังหวัดโค้ด22 อำเภอ2201–2210 กำหนดชื่อจากPhase11B ไม่สร้างสวน/พิกัด GISไม่เข้าสูตร. Sampling uncertainty/CIไม่มีฐานประเมินจาก2ปี ไม่สร้างขึ้น. ห้ามใช้ผลประมาณฝึกMLหรืออนุมานดิน/อุณหภูมิเชิงสาเหตุ.

## ทำซ้ำ
จาก C:/project ใช้ Python ที่มีopenpyxl/pypdf:
`python -m tools.build_district_spatial_allocation --output-name allocation_reproduce_new_id`
เขียนได้เฉพาะโฟลเดอร์ใหม่ ปฏิเสธoverwrite. ไม่ติดตั้งDependencies ไม่ฝึกML. Source ManifestมีHashes/Anchors; QC ReportแยกFixed arithmetic passกับArea blocked. ตรวจProtected1288รายการกับHashรายงานเดิมก่อน/หลัง. ใช้PonytailแยกScriptเฉพาะงาน ไม่สร้างระบบใหม่.
''')
    assert all(sha(ROOT/k)==v for k,v in before.items())
    assert all(sha(ROOT/k)==v for k,v in protected.items())
    save(out/'qc_report.json',dict(fixed_status='arithmetic_pass_historical_accuracy_unverified',area_adjusted_status='blocked_unresolved_source_conflict',checks=checks,source_hash_mismatches=[],protected_files_checked=len(protected),protected_hash_mismatches=[],production_before=production,production_after=sha(MODEL),oae_reconciliation=reconciliation,conflicts=conflicts,estimated_numeric_rows=70,blocked_rows=70,reported_rows=20,publication_vintage='unverified',public_release='HOLD'))
    print(json.dumps(dict(output=str(out),validation=summaries,conflicts=conflicts,checks=len(checks)),ensure_ascii=True))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output-name',default='allocation_20261012_v1');args=parser.parse_args()
    assert Path(args.output_name).name==args.output_name and args.output_name not in ('.','..')
    build(ROOT/'research_data/internal_research/district_spatial_allocation'/args.output_name)
