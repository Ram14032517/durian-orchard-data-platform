"""Offline Chanthaburi analysis. Only writes a new internal output directory."""
import argparse
import csv
import json
import platform
from pathlib import Path
from statistics import mean, correlation
import numpy as np
from tools.build_chanthaburi_advisor_eda import ROOT, P9, P10, P11, MODEL, sha, read, rows, seal_check, svg_plot, table, e

PROTOCOL = dict(issue_month_day='03-01', weather_months=list(range(1,13)),
    weather_lag_years=1, yield_lag_years=2, label_embargo_years=2,
    features=['yield_lag2_kg_rai','t2m_lag1_c','rain_lag1_mm','rh_lag1_pct'],
    first_training_year=2008, inner_years=list(range(2015,2019)),
    test_years=list(range(2020,2026)), alpha_grid=[0.1,1.0,10.0],
    scaling='training-only population mean/std', intercept='unpenalized',
    gate=dict(inner_mae_gain_min=.05,test_mae_gain_min=.10,test_rmse_not_worse=True,paired_wins_min=4),
    prediction_policy='unclamped; negative prediction invalidates scenario gate',
    mode='exploratory; observation cutoff verified, publication vintage unverified',
    scenario='Only if gate passes; lagged annual temperature sensitivity, never causal future warming',
    holdout_warning='2020-2025 already inspected in earlier project phases; not independent confirmation')

def save(p,v):
    with p.open('x',encoding='utf-8') as f: json.dump(v,f,ensure_ascii=False,indent=2,allow_nan=False)

def textfile(p,v):
    with p.open('x',encoding='utf-8') as f: f.write(v)

def writecsv(p,data):
    assert data
    with p.open('x',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(data[0])); w.writeheader(); w.writerows(data)

def ridge(X,y,q,alpha):
    X=np.array(X,dtype=float); y=np.array(y,dtype=float); q=np.array(q,dtype=float)
    assert X.ndim==2 and len(X)==len(y) and len(X)>=3
    assert np.isfinite(X).all() and np.isfinite(y).all() and np.isfinite(q).all()
    mu=X.mean(axis=0); sd=X.std(axis=0); sd=np.where(sd==0,1,sd)
    Z=(X-mu)/sd; ym=y.mean()
    beta=np.linalg.solve(Z.T@Z+alpha*np.eye(X.shape[1]),Z.T@(y-ym))
    pred=float(ym+(q-mu)/sd@beta)
    return pred,dict(mean=mu.tolist(),std=sd.tolist(),beta_standardized=beta.tolist(),intercept_target_mean=float(ym))

def metrics(rr):
    err=np.array([r['error_kg_rai'] for r in rr])
    return dict(n=len(rr),mae=float(np.abs(err).mean()),rmse=float(np.sqrt((err**2).mean())),bias=float(err.mean()))

def selftests():
    # Analytic two-feature reference, intercept and train-only standardization.
    X=np.array([[0,2],[1,4],[2,6],[3,8]],dtype=float); y=np.array([2,4,6,8.])
    p,a=ridge(X,y,[4,10],1)
    Z=(X-X.mean(0))/X.std(0); q=(np.array([4,10])-X.mean(0))/X.std(0)
    expected=float(y.mean()+q@np.linalg.inv(Z.T@Z+np.eye(2))@Z.T@(y-y.mean()))
    assert abs(p-expected)<1e-10
    p2,_=ridge(X,y+100,[4,10],1); assert abs(p2-p-100)<1e-10
    _,b=ridge(X,y,[100,200],1); assert a['mean']==b['mean'] and a['std']==b['std']
    assert abs(ridge([[1],[1],[1]],[2,3,4],[1],1)[0]-3)<1e-12
    try: ridge([[float('nan')],[1],[2]],[1,2,3],[1],1)
    except AssertionError: pass
    else: raise AssertionError('Missing features accepted')
    return ['ridge independent inverse reference','intercept translation','query cannot affect scaler','constant feature','missing feature fail-closed']

def chart(xx,yy,title,xlabel,ylabel,color,line=True):
    lo,hi=min(yy),max(yy); pad=max((hi-lo)*.15,1 if ylabel!='°C' else .1)
    return svg_plot(xx,yy,title,xlabel,ylabel,(min(xx),max(xx)),(max(0,lo-pad),hi+pad),color,line=line,years=list(range(2006,2026)))

def build(out):
    assert not out.exists(),'Refuse overwrite'
    checks=selftests()
    protected=read(ROOT/'research_data/internal_research/chanthaburi_district_yield/phase11e_20261012_v1/preflight_protected_hashes.json')
    assert all(sha(ROOT/k)==v for k,v in protected.items()),'Protected input mismatch'
    for p in (P9,P10,P11): seal_check(p)
    oldeda=ROOT/'research_data/internal_research/chanthaburi_eda/eda_20261012_v1'
    seal_check(oldeda)
    production=sha(MODEL)
    assert production=='b4b9e0985e6ba5e6ee371d8988cad5e65dee031b345eb0cf13a8695f91658680'
    qc=read(P10/'qc_summary.json'); assert qc['status']=='pass'
    annual=sorted([r for r in rows(P9/'province_year.csv') if r['province_code']=='22'],key=lambda r:int(r['year_ce']))
    assert [int(r['year_ce']) for r in annual]==list(range(2006,2026))
    monthly=[r for r in rows(P9/'weather_profile_month.csv') if r['province_code']=='22']
    assert len(monthly)==240 and len({(r['year_ce'],r['month']) for r in monthly})==240
    for r in annual:
        mm=[m for m in monthly if m['year_ce']==r['year_ce']]
        assert sorted(int(m['month']) for m in mm)==list(range(1,13))
        days=sum(int(m['days']) for m in mm)
        for raw,key in [('T2M','weather_t2m_mean_c'),('RH2M','weather_rh2m_mean_pct')]:
            assert abs(sum(float(m[raw])*int(m['days']) for m in mm)/days-float(r[key]))<1e-6
            assert all(float(m[raw+'_area_coverage'])>.999999 for m in mm)
        assert abs(sum(float(m['PRECTOTCORR'])*int(m['days']) for m in mm)-float(r['weather_precip_total_mm']))<1e-6
    # Independent raw OAE verification, retained cell anchors from earlier read-only EDA.
    from openpyxl import load_workbook
    wbpath=ROOT/'research_data/external/oae_2026-09-17/durian_province.xlsx'
    wb=load_workbook(wbpath,read_only=True,data_only=True)
    mapping={'ผลผลิตต่อเนื้อที่ให้ผล':'yield_kg_per_rai','เนื้อที่ให้ผล':'bearing_area_rai','เนื้อที่ยืนต้น':'planted_area_rai','ผลผลิต':'production_tonnes'}
    cells=[]
    for rownum,r in enumerate(wb['durian'].iter_rows(min_row=2,values_only=True),2):
        year,crop,province,commodity,subcom,item,value,unit,date=r
        if province=='จันทบุรี' and commodity=='ทุเรียน' and item in mapping and 2549<=int(year)<=2568:
            target=annual[int(year)-2549]
            assert float(target[mapping[item]])==float(value)
            assert unit=={'yield_kg_per_rai':'กิโลกรัมต่อไร่','production_tonnes':'ตัน','bearing_area_rai':'ไร่','planted_area_rai':'ไร่'}[mapping[item]]
            cells.append(dict(year_ce=int(year)-543,field=mapping[item],sheet='durian',cell=f'G{rownum}',data_date=str(date)))
    wb.close(); assert len(cells)==80
    district=read(P11/'district_year.json')['rows']
    assert len(district)==20 and {(r['district_code'],r['year_be']) for r in district}=={(str(c),y) for c in range(2201,2211) for y in (2566,2567)}
    soil=read(P10/'dataset.json')['tables']; ss=soil['soil_composition']; den=ss[0]['denominator_pure_rai']
    unc=sum(r['area_rai'] for r in soil['uncovered'])
    assert abs(sum(r['area_rai'] for r in ss)+unc-den)<=.1
    assert abs(sum(r['share_of_all_pure_pct'] for r in ss)+100*unc/den-100)<=100*.1/den
    checks+=['20 annual years complete','240 unique months','temperature/rain/humidity annual reconciliation','80 raw OAE cells and units','20 district keys','sealed accepted soil reconciliation']
    inputs=[P9/n for n in ['province_year.csv','weather_profile_month.csv','sampling_points.csv','response_signatures.csv','soil_composition.csv','source_manifest.json','output_manifest.json']]+[P10/n for n in ['dataset.json','qc_summary.json','release_metadata.json','source_manifest.json','output_manifest.json']]+[P11/n for n in ['district_year.json','source_manifest.json','field_lineage.json','output_manifest.json']]+[wbpath,oldeda/'output_manifest.json',Path(__file__),ROOT/'tools/build_chanthaburi_advisor_eda.py']
    inputhash={str(p.relative_to(ROOT)):sha(p) for p in inputs}
    out.mkdir(parents=True)
    # Protocol is persisted before any analysis fit, never modified after scores.
    save(out/'protocol.json',PROTOCOL)
    save(out/'input_manifest.json',dict(input_hashes=inputhash,production_sha256=production,protected_files=len(protected),runtime=platform.python_version(),numpy=np.__version__))
    data=[]; byyear={int(r['year_ce']):r for r in annual}
    for year,r in byyear.items():
        d=dict(province_code='22',year_ce=year,year_be=year+543,yield_kg_rai=float(r['yield_kg_per_rai']),temperature_c=float(r['weather_t2m_mean_c']),rain_mm=float(r['weather_precip_total_mm']),humidity_pct=float(r['weather_rh2m_mean_pct']),bearing_area_rai=float(r['bearing_area_rai']),production_tonnes=float(r['production_tonnes']),issue_date=f'{year}-03-01',feature_observation_end=f'{year-1}-12-31' if year>=2008 else None,
            yield_lag2_kg_rai=float(byyear[year-2]['yield_kg_per_rai']) if year>=2008 else None,
            t2m_lag1_c=float(byyear[year-1]['weather_t2m_mean_c']) if year>=2008 else None,
            rain_lag1_mm=float(byyear[year-1]['weather_precip_total_mm']) if year>=2008 else None,
            rh_lag1_pct=float(byyear[year-1]['weather_rh2m_mean_pct']) if year>=2008 else None,
            split='test' if year in PROTOCOL['test_years'] else 'inner' if year in PROTOCOL['inner_years'] else 'historical',
            availability_status='unverified_publication_vintage',weather_footprint_year_be=2568)
        data.append(d)
    usable={d['year_ce']:d for d in data if d['yield_lag2_kg_rai'] is not None}
    def fold(year,model,alpha,split):
        train=[d for y,d in usable.items() if y<=year-2]
        q=usable[year]; assert max(d['year_ce'] for d in train)<=year-2
        assert q['feature_observation_end']<q['issue_date'] and year not in [d['year_ce'] for d in train]
        assert all(d['year_ce']-1<=year-3 for d in train)
        if model=='ridge': pred,fit=ridge([[d[k] for k in PROTOCOL['features']] for d in train],[d['yield_kg_rai'] for d in train],[q[k] for k in PROTOCOL['features']],alpha)
        else: pred=q['yield_lag2_kg_rai'] if model=='lag2_baseline' else mean(d['yield_kg_rai'] for d in train); fit=None
        return dict(record_type='prediction',split=split,model=model,alpha=alpha,year_ce=year,n_train=len(train),fit_label_year_max=max(d['year_ce'] for d in train),issue_date=q['issue_date'],feature_observation_end=q['feature_observation_end'],actual_kg_rai=q['yield_kg_rai'],predicted_kg_rai=pred,error_kg_rai=pred-q['yield_kg_rai'],absolute_error_kg_rai=abs(pred-q['yield_kg_rai']),availability_status='unverified_publication_vintage'),fit
    inner=[]
    for alpha in PROTOCOL['alpha_grid']:
        for year in PROTOCOL['inner_years']: inner.append(fold(year,'ridge',alpha,'inner')[0])
    scores=[dict(alpha=a,**metrics([r for r in inner if r['alpha']==a])) for a in PROTOCOL['alpha_grid']]
    selected=min(scores,key=lambda r:(r['mae'],r['alpha']))['alpha']
    rr=[r for r in inner if r['alpha']==selected]
    for model in ('lag2_baseline','training_mean'):
        rr += [fold(year,model,None,'inner')[0] for year in PROTOCOL['inner_years']]
    fits=[]
    for model in ('lag2_baseline','training_mean','ridge'):
        for year in PROTOCOL['test_years']:
            record,fit=fold(year,model,selected if model=='ridge' else None,'test'); rr.append(record)
            if fit: fits.append(dict(year_ce=year,alpha=selected,**fit))
    summary={split:{m:metrics([r for r in rr if r['split']==split and r['model']==m]) for m in ('lag2_baseline','training_mean','ridge')} for split in ('inner','test')}
    gains=[next(r['absolute_error_kg_rai'] for r in rr if r['model']=='lag2_baseline' and r['split']=='test' and r['year_ce']==y)-next(r['absolute_error_kg_rai'] for r in rr if r['model']=='ridge' and r['split']=='test' and r['year_ce']==y) for y in PROTOCOL['test_years']]
    wins=sum(g>1e-9 for g in gains); ties=sum(abs(g)<=1e-9 for g in gains)
    igain=1-summary['inner']['ridge']['mae']/summary['inner']['lag2_baseline']['mae']; tgain=1-summary['test']['ridge']['mae']/summary['test']['lag2_baseline']['mae']
    gate=dict(inner_gain=igain,test_gain=tgain,inner_pass=igain>=.05,test_mae_pass=tgain>=.10,test_rmse_pass=summary['test']['ridge']['rmse']<=summary['test']['lag2_baseline']['rmse'],wins_pass=wins>=4,nonnegative_pass=all(r['predicted_kg_rai']>=0 for r in rr),wins=wins,ties=ties,losses=6-wins-ties)
    gate['passed']=all(gate[k] for k in ('inner_pass','test_mae_pass','test_rmse_pass','wins_pass','nonnegative_pass'))
    # Outer target mutation must not change fit or predictions in the same fold.
    q=usable[2025]; before_pred=fold(2025,'ridge',selected,'test')[0]['predicted_kg_rai']; old=q['yield_kg_rai']
    q['yield_kg_rai']=old+100000; after_pred=fold(2025,'ridge',selected,'test')[0]['predicted_kg_rai']; q['yield_kg_rai']=old
    assert before_pred==after_pred
    assert len({(r['split'],r['model'],r['year_ce']) for r in rr})==30
    assert all(r['year_ce'] not in PROTOCOL['test_years'] for r in inner)
    checks+=['all folds observation/label cutoff','outer target mutation cannot alter prediction','30 unique selected-arm predictions','alpha selection excludes outer years']
    # Scenario is additionally not supported if feature perturbations leave training support.
    scenario=[]
    if gate['passed']:
        train=[d for y,d in usable.items() if y<=2023]; q=usable[2025]
        tmin=min(d['t2m_lag1_c'] for d in train); tmax=max(d['t2m_lag1_c'] for d in train)
        for delta in (-2,-1,1,2):
            valid=tmin<=q['t2m_lag1_c']+delta<=tmax
            scenario.append(dict(delta_c=delta,status='unsupported_extrapolation' if not valid else 'descriptive_feature_sensitivity_only',prediction_kg_rai=None))
        # No unsupported values or calibrated interval is fabricated from six reused test years.
    scenario_status='not_trustworthy_no_predictions_issued' if not gate['passed'] or any(r['status']=='unsupported_extrapolation' for r in scenario) else 'uncertainty_not_calibrated_no_scenario_issued'
    years=[d['year_ce'] for d in data]; yieldvals=[d['yield_kg_rai'] for d in data]
    correlations={k:correlation([d[k] for d in data],yieldvals) for k in ('temperature_c','rain_mm','humidity_pct')}
    summaryall=dict(metrics=summary,alpha_grid_scores=scores,selected_alpha=selected,gate=gate,scenario_status=scenario_status,scenario_support_checks=scenario,pearson=correlations,checks=checks,source_cells=cells,soil_denominator_rai=den,soil_uncovered_rai=unc,paired_gains=gains)
    writecsv(out/'analysis_results.csv',data); writecsv(out/'model_evaluation.csv',rr); writecsv(out/'inner_grid_evaluation.csv',inner)
    save(out/'analysis_summary.json',summaryall); save(out/'ridge_fold_parameters.json',fits)
    # GIS/district remain separate grains, never model features.
    writecsv(out/'tha_mai_soil_summary.csv',[dict(soil_unit=r['soil_unit'],area_rai=r['area_rai'],share_pct=r['share_of_all_pure_pct'],denominator_rai=den) for r in ss])
    writecsv(out/'district_case_study.csv',[dict(district_code=r['district_code'],district_name=r['district_name'],year_be=r['year_be'],reported_yield_kg_rai=r['reported_yield_kg_rai'],production_tonnes=r['production_tonnes'],bearing_area_rai=r['bearing_area_rai'],status='source_transcription_only_finality_unverified') for r in district])
    def page(title,body): return f'<section class="page"><header>Selected Topics in AI <span>จันทบุรี / Internal Research</span></header><h1>{title}</h1>{body}<footer>Public Release HOLD | ไม่รับรอง Historical PIT <span>{len(pages)+1} / 10</span></footer></section>'
    pages=[]
    f=lambda x:f'{x:,.2f}'
    drop=100*(byyear[2024] and float(byyear[2024]['yield_kg_per_rai'])/float(byyear[2023]['yield_kg_per_rai'])-1)
    top=sorted(ss,key=lambda r:-r['area_rai'])
    findings=[f'Yield จังหวัดสูงสุดปี 2022 = 2,194 กก./ไร่ ต่ำสุดปี 2016 = 1,094 กก./ไร่ ไม่เพิ่มสม่ำเสมอทุกปี',f'ปี 2024 Yield ลด {abs(drop):.1f}% จากปี 2023 และอุณหภูมิเฉลี่ย {float(byyear[2024]["weather_t2m_mean_c"]):.2f} °C แต่ไม่พิสูจน์เหตุและผล',f'ตารางรายอำเภอปี 2567 มี Yield ต่ำกว่าปี 2566 ทั้ง 10 อำเภอ สถานะ Final Actual ยังไม่ยืนยัน',f'สามหน่วยแผนที่ดินใหญ่สุดใน Pure A403 ท่าใหม่รวม {sum(r["share_of_all_pure_pct"] for r in top[:3]):.2f}% ไม่ใช้แทนดินทุกอำเภอ',f'Ridge Test MAE {f(summary["test"]["ridge"]["mae"])} เทียบ Baseline {f(summary["test"]["lag2_baseline"]["mae"])} กก./ไร่; ผ่านเกณฑ์ทั้งหมด: {"ใช่" if gate["passed"] else "ไม่"}']
    pages.append(page('ทุเรียนจันทบุรี: ข้อมูลบอกอะไรได้จริง',f'<p class="lead">ผลวิเคราะห์ข้อมูลเปิดที่มีอยู่ สำหรับหารืออาจารย์ ไม่ใช่ผลยืนยันความแม่นของระบบ Production</p><p>12 ตุลาคม 2569 | จังหวัด–ปี 2006–2025 | กรณีรายอำเภอ 2566–2567</p><h2>ข้อค้นพบหลัก</h2><ol>'+''.join('<li>'+e(v)+'</li>' for v in findings)+'</ol><div class="notice"><h2>ข้อสรุปเรื่อง Scenario</h2><p>ยังไม่ออกค่าพยากรณ์สถานการณ์อุณหภูมิ ±1°C และ ±2°C เพราะเกณฑ์ความน่าเชื่อถือ/ขอบเขตข้อมูลยังไม่ครบ ไม่สร้าง Confidence Interval ที่ไม่มีหลักฐาน</p></div><h2>ขอบเขตการนำไปใช้</h2><p>ใช้เปิดประเด็นเรื่องความผันผวนของผลผลิต ความสัมพันธ์ของอากาศ และข้อจำกัดของหน่วยข้อมูล ไม่ใช้แนะนำผลผลิตรายสวนหรือผลกระทบจากโลกร้อนเชิงสาเหตุ</p>'))
    pages.append(page('ข้อมูลและหน่วยวิเคราะห์',table(['ข้อมูล','Coverage จริง','หน่วย / ข้อจำกัด'],[['OAE จังหวัด','20 ปี 2006–2025','22–ปี; Yield ที่รายงาน'],['POWER profile','240 เดือน, 16 จุดร้องขอ','ถ่วงพื้นที่ A403 รุ่น 2568'],['A403 จันทบุรี','8,512 Pure polygons','แผนที่การใช้ที่ดิน ไม่ใช่ทะเบียนสวน'],['ดินท่าใหม่','1,504 Pure fragments / 43 รหัส','ผ่าน Policy v2; GIS ยัง HOLD'],['Yield รายอำเภอ','10 อำเภอ × 2 ปี','ถอดต้นฉบับ ไม่รับรอง Final Actual']])+ '<h2>การรวมอากาศ</h2><p>ช่วงเดือนมกราคม–ธันวาคมทุกปี T2M (°C) และ RH2M (%) เฉลี่ยถ่วงจำนวนวัน ฝน PRECTOTCORR (มม./วัน) คูณวันแล้วรวมทั้งปี ตรวจครบ 12 เดือน ไม่มีการเติม Missing หรือเปลี่ยนหน่วย</p><h2>ข้อจำกัดเชิงพื้นที่</h2><p>POWER เป็นอากาศกริด ไม่ใช่เซนเซอร์สวน แหล่งเดิมระบุ Meteorology ประมาณ 0.5°×0.625° ระดับหลายสิบกิโลเมตร 16 จุดให้เพียง 6 Meteorology response signatures (11 เมื่อรวม Solar) ไม่เท่ากับจำนวนสถานีหรือจำนวน Native Grid ที่ยืนยันแล้ว</p><p>น้ำหนักและ footprint A403 ปี 2568 ถูกใช้ย้อนหลัง เป็น Retrospective Fixed Footprint ไม่ใช่ประวัติสวนจริงทุกปี Phase 12C ไม่พบชุด Integrated Dataset ที่พร้อมใช้ จึงไม่รวมทะเบียนกลุ่มแปลงใหญ่หรือหมุดสวน</p><h2>คุณภาพและความพร้อม</h2><p>Province-year ตรวจ 80 ค่า OAE กับ Workbook ต้นทางและ 240 เดือนอากาศ ค่าดินใช้ Accepted Proposal เดิม ไม่มี Overlay ใหม่ License การเผยแพร่ Derived LDD ยังไม่ยืนยัน Public HOLD คงเดิม</p>'))
    pages.append(page('แนวโน้ม Yield และอุณหภูมิ', '<h2>ผลผลิตต่อพื้นที่ให้ผล (กก./ไร่)</h2>'+chart(years,yieldvals,'Yield','ปี ค.ศ.','กก./ไร่','#2563eb')+'<h2>อุณหภูมิเฉลี่ยทั้งปี (°C)</h2>'+chart(years,[d['temperature_c'] for d in data],'T2M','ปี ค.ศ.','°C','#087f8c')+'<p>กราฟอุณหภูมิไม่เริ่มแกนที่ศูนย์ แยกแกนเพื่อหลีกเลี่ยงภาพชวนให้คิดว่าเส้นสองชนิดเคลื่อนเท่ากัน ช่วงเดือนของกราฟย้อนหลังเป็นปีเดียวกับ Yield ใช้เพื่อพรรณนาเท่านั้น อากาศครบปีเป้าหมายไม่เข้าโมเดล</p>'))
    pages.append(page('แนวโน้มฝนและความชื้น','<h2>ฝนรวมมกราคม–ธันวาคม (มม./ปี)</h2>'+chart(years,[d['rain_mm'] for d in data],'Rain','ปี ค.ศ.','มม./ปี','#2563eb')+'<h2>ความชื้นสัมพัทธ์เฉลี่ยทั้งปี (%)</h2>'+chart(years,[d['humidity_pct'] for d in data],'RH2M','ปี ค.ศ.','%','#087f8c')+'<p>ค่าเฉลี่ยรายปีอาจกลบฝนทิ้งช่วง ความร้อนสุดขั้ว และช่วงออกดอก/ติดผล ความชื้น 2 เมตรไม่ใช่ความชื้นดิน ข้อมูลนี้ยังไม่ระบุการให้น้ำ อายุสวน พันธุ์ หรือความเสียหาย</p>'))
    scatter=''
    for key,label,unit in [('temperature_c','อุณหภูมิ','°C'),('rain_mm','ฝน','มม./ปี'),('humidity_pct','ความชื้น','%')]:
        vals=[d[key] for d in data]; lo,hi=min(vals),max(vals); pad=(hi-lo)*.1
        scatter+=f'<h2>{label} กับ Yield: Pearson r = {correlations[key]:.3f}</h2><div class="scatter">'+svg_plot(vals,yieldvals,label,f'{label} ({unit})','Yield กก./ไร่',(lo-pad,hi+pad),(800,2400),'#087f8c',years=years)+'</div>'
    pages.append(page('ความสัมพันธ์ย้อนหลัง ไม่ใช่เหตุและผล',scatter+'<p class="small">n = 20 จังหวัด–ปี หนึ่งจุดเป็นหนึ่งปี ไม่ใช่สวน 20 แห่ง และปีติดกันไม่จำเป็นต้องอิสระ r ใช้ค่าดิบ ไม่ควบคุม Trend หรือ Confounders ไม่รายงาน p-value/นัยสำคัญ และไม่ใช้ r เลือก Features หลังเห็นผล</p>'))
    pages.append(page('การทดลองโมเดลโดยไม่ใช้อากาศอนาคต',table(['ขั้นตอน','กติกาที่ล็อกก่อน Fit'],[['วันจำลองพยากรณ์','1 มีนาคมของปี Y (กฎใหม่สำหรับชั้นเรียน)'],['Features','Yield Y−2; T2M / Rain / RH ปี Y−1'],['ช่วงอากาศ','ม.ค.–ธ.ค. Y−1 สิ้นสุดก่อน Issue Date'],['Training Labels','ปีเป้าหมาย ≤ Y−2 ทุก Fold'],['เลือก Alpha','0.1 / 1 / 10 จาก Inner ปี 2015–2018 เท่านั้น'],['Test','2020–2025 รวม 6 ปี; Expanding Training'],['Baseline','Yield Y−2 และค่าเฉลี่ย Training'],['Preprocessing','Mean/Std เฉพาะ Training; ไม่มีสุ่มแถว']])+ '<h2>เหตุผลของช่องว่างข้อมูล</h2><p>ไม่ใช้ Yield ปีก่อนทันที เพราะวันเผยแพร่ OAE ยังไม่ยืนยัน เลือก Lag 2 เป็นสมมติฐานเผื่อเวลา ไม่ใช่หลักฐาน PIT สำหรับ 2020 จึง Fit เฉพาะ 2008–2018 และใช้ Weather ปี 2019 ก่อนวันจำลองพยากรณ์</p><h2>เกณฑ์ Scenario ที่ประกาศล่วงหน้า</h2><p>Inner MAE ดีขึ้น ≥5%, Test MAE ดีขึ้น ≥10% เทียบ Lag 2, Test RMSE ไม่แย่ลง, ชนะอย่างน้อย 4/6 ปี และไม่มีค่าติดลบ/Integrity Failure ต้องอยู่ในช่วงข้อมูล Training จึงพิจารณาความไวต่ออุณหภูมิได้</p><h2>สิ่งที่ยังไม่ป้องกันได้ด้วยการแบ่งปี</h2><p>OAE และ POWER เป็นฉบับย้อนหลังที่มีความเสี่ยง Revised Data วันที่เข้าถึง/ดาวน์โหลดไม่ใช่ Publication Date แผนที่ 2568 เป็นข้อมูลภายหลัง และปี Test เคยถูกตรวจใน Phase เดิมแล้ว ผลนี้เป็น Exploratory ไม่ใช่ Holdout อิสระหรือ Historical PIT Certified ไม่เปรียบเทียบกับ MAE 204.85 คนละ Protocol</p>'))
    mt=[]
    for split in ('inner','test'):
        for model,v in summary[split].items(): mt.append([split,model,v['n'],f(v['mae']),f(v['rmse']),f(v['bias'])])
    predtable=[[y,f(usable[y]['yield_kg_rai'])]+[f(next(r['predicted_kg_rai'] for r in rr if r['split']=='test' and r['year_ce']==y and r['model']==m)) for m in ('lag2_baseline','ridge')] for y in PROTOCOL['test_years']]
    pages.append(page('ผลทดสอบและการตัดสินใจ Scenario',f'<p>Selected Alpha = {selected:g} ตรึงจาก Inner ไม่เลือกจาก Test Bias = Prediction − Actual ค่าลบหมายถึงทำนายต่ำ</p>'+table(['ชุด','โมเดล','n','MAE','RMSE','Bias'],mt)+'<h2>Test รายปี (กก./ไร่)</h2>'+table(['ปี ค.ศ.','จริง','Baseline Lag 2','Ridge'],predtable)+f'<p>Paired Ridge เทียบ Lag 2: ชนะ {wins} / เสมอ {ties} / แพ้ {6-wins-ties} ปี Test MAE เปลี่ยน {tgain*100:.2f}% ในทิศทางดีขึ้น Inner MAE เปลี่ยน {igain*100:.2f}%</p><p>Gate: Inner {gate["inner_pass"]}, Test MAE {gate["test_mae_pass"]}, RMSE {gate["test_rmse_pass"]}, Wins {gate["wins_pass"]}. ผ่านทั้งหมด = {gate["passed"]}</p><div class="notice"><p>Scenario ยังไม่น่าเชื่อถือ จึงไม่สร้างผล ±1/±2°C หรือ Confidence Interval อากาศ Feature เป็นของปีก่อน ไม่ใช่อุณหภูมิฤดูอนาคต สัมประสิทธิ์ Ridge ไม่ใช่ Causal Effect และ 6 Test ปีไม่พอรับรองช่วงความไม่แน่นอนที่สอบเทียบแล้ว</p></div>'))
    soilrows=[[r['soil_unit'],f(r['area_rai']),f'{r["share_of_all_pure_pct"]:.4f}'] for r in top[:8]]
    others=sum(r['area_rai'] for r in top[8:] if r['soil_unit'] not in ('SC','W'))
    for code in ('SC','W'):
        r=next(r for r in ss if r['soil_unit']==code)
        if code not in [r[0] for r in soilrows]: soilrows.append([code,f(r['area_rai']),f'{r["share_of_all_pure_pct"]:.6f}'])
    soilrows+=[['Other (แสดงผล)',f(others),f'{others/den*100:.4f}'],['Uncovered',f'{unc:.3e}',f'{unc/den*100:.3e}']]
    pages.append(page('A403 และดิน: หน่วยพื้นที่ที่ต้องแยก',f'<p>จันทบุรีมี Pure A403 8,512 Polygons แผนที่การใช้ที่ดิน พ.ศ.2568 พื้นที่ {539971.790558:,.2f} ไร่ ไม่ใช่พื้นที่เหมาะสม ไม่ใช่สวนรายเจ้าของ Soil Production Year 2561 ไม่ใช่วันที่สำรวจทุก Polygon</p><h2>ดินเฉพาะ Pure A403 ท่าใหม่</h2><p>ตัวหาร Pure A403 ทั้งอำเภอ {den:,.2f} ไร่ รวม SC/W/Uncovered ไม่รวม Mixed-use ที่ไม่ทราบส่วนทุเรียน CRS EPSG:32647 พื้นที่ไร่ = m² / 1,600</p>'+table(['หน่วยแผนที่ดิน','พื้นที่ไร่','% ตัวหาร Pure'],soilrows)+'<p>Other รวมเฉพาะเพื่อแสดงผล ตารางต้นฉบับ 43 รหัสไม่เปลี่ยน SC เป็นพื้นที่ลาดชันเชิงซ้อน W เป็นน้ำ อยู่ใน Covered แล้ว ไม่บวกพื้นที่ซ้ำ Mixed-use เก็บแยก ไม่จัดสรรเป็นทุเรียนโดยเดา</p><h2>Geometry QC</h2><p>Primary Policy v2 ผ่าน Pure Fragments 1,504 และ Parent Partitions 3,025 ชุดเดิมมี GEOS 3.13.1 Two-stage Union Warning เติมรู 14.2778 ไร่ใน Diagnostic Path ไม่ใช่พื้นที่ดินที่แก้ใหม่ Tolerance เดิม 0.1 ไร่ ไม่มีการ Union/Overlay ในงานนี้</p><p>สัดส่วนดินท่าใหม่ไม่ใช่ตัวแทนอีก 9 อำเภอ ไม่มี Yield ต่อ Polygon และไม่มีความสัมพันธ์ดิน–Yield รายสวนที่ยืนยันได้</p>'))
    dt=[]
    for c in range(2201,2211):
        aa=next(r for r in district if r['district_code']==str(c) and r['year_be']==2566); bb=next(r for r in district if r['district_code']==str(c) and r['year_be']==2567)
        dt.append([c,aa['district_name'],aa['reported_yield_kg_rai'],bb['reported_yield_kg_rai'],f((bb['reported_yield_kg_rai']/aa['reported_yield_kg_rai']-1)*100)])
    pages.append(page('กรณีศึกษา Yield รายอำเภอ',table(['รหัส','อำเภอ','2566 กก./ไร่','2567 กก./ไร่','เปลี่ยน %'],dt)+'<h2>จุดที่ควรถามต้นทาง</h2><p>ปี 2566 ท่าใหม่มี Yield สูงสุด 2,220 กก./ไร่ ปี 2567 เขาคิชฌกูฏสูงสุด 1,625 กก./ไร่ สอยดาวต่ำสุดทั้งสองปี 905 และ 724 กก./ไร่ ทุกอำเภอลดลง แต่ยังอธิบายไม่ได้ว่าเป็นฤดูอากาศ การจัดการหรือสถานะประมาณการ</p><p>Source: PDF กรมส่งเสริมการเกษตร หน้า 5 ตารางที่ 2 อ้างสำนักงานเศรษฐกิจการเกษตรที่ 6 Phase 11B v2 รับรองเฉพาะการถอดตาราง วันที่เผยแพร่ นิยามปี และ Final Actual ยัง Unverified ไม่เอาค่าแต่ละอำเภอไปแจกให้สวน</p><p>สองปีไม่เพียงพอแยกผลดิน/ภูมิอากาศ/การจัดการอย่างเสถียร ดินพร้อมเฉพาะท่าใหม่ และ District Weather Exposure ยังไม่พร้อม ไม่มีการ Join Soil–Weather–Yield รายอำเภอในงานนี้</p><h2>ข้อมูลจาก Phase 11C–11F</h2><p>ข้อมูลพื้นที่ยืนต้นและ Source Leads ไม่ใช่ผลผลิตย้อนหลังที่ยืนยันได้ จึงไม่เติมปี 2014–2022 ร่างคำขอข้อมูลยังไม่ใช่ Dataset และไม่ได้ติดต่อหน่วยงาน</p>'))
    pages.append(page('ข้อจำกัด แหล่งข้อมูล และการใช้ต่อ', '<h2>ประเด็นให้อาจารย์ตัดสินใจ</h2><ol><li>เน้นโจทย์ระดับจังหวัดเชิงพรรณนา หรือระดับอำเภอที่ต้องมี Yield หลายปีและข้อมูลอากาศตรงหน่วยพื้นที่?</li><li>ควรกำหนดช่วงอากาศตามฤดูออกดอก/ติดผลด้วยหลักฐานใด แทนค่าเฉลี่ยทั้งปี โดยไม่เลือกเดือนเพื่อไล่คะแนน?</li><li>ยอมรับ Fixed Footprint 2568 เป็น Retrospective Proxy หรือจำเป็นต้องมีประวัติพื้นที่ปลูก?</li><li>เกณฑ์ข้อมูล/สิทธิ์ใดต้องยืนยันก่อนใช้คำว่า Forecast หรือเผยแพร่ Derived GIS?</li></ol><h2>แหล่งข้อมูลภายในที่ใช้</h2><p>[1] Phase 9D v1: province_year / weather_profile_month / Source Manifest เชื่อม OAE จังหวัดกับ NASA POWER spatial profile<br>[2] OAE durian_province.xlsx (oae_2026-09-17), Sheet durian, 80 Cell Anchors ใน analysis_summary.json<br>[3] Phase 10F Internal Accepted Proposal v1: dataset / QC Summary / Release Metadata ใช้เฉพาะสัดส่วนดิน<br>[4] Phase 11B v2: district_year / Source Manifest / Field Lineage, PDF หน้า 5 Table 2</p><h2>การทำซ้ำและการปกป้องข้อมูล</h2><p>อ่าน Source แบบ Read-only ตรวจ Seals/Hashes และ Production SHA256 ก่อน–หลัง บันทึก Protocol ก่อน Fit มี Train-only Scaling, Fold Cutoffs และ Ridge Coefficients รายปี การตรวจเฉพาะงานอยู่ใน qc_report.json ไม่มี Repo-wide Tests ที่ไม่เกี่ยวข้อง ไม่มี Network หรือ API Changes</p><p>ไฟล์ analysis_results.csv, model_evaluation.csv, data_dictionary.md, source_manifest.md และ README_TH.md อยู่ในโฟลเดอร์เดียวกัน พร้อมโปรแกรม reproduce.py ใช้ NumPy/OpenPyXL และ helper เดิม การรันซ้ำต้องใช้โฟลเดอร์ใหม่ ห้ามเขียนทับ</p><div class="notice"><p>สิทธิ์เผยแพร่ Derived LDD ยัง Pending แม้เป็นข้อมูลบนเว็บไซต์รัฐ งานนี้ Internal Only / Public HOLD ไม่อัปโหลด Geometry หรือ Dataset ผลการแบ่งเวลานี้ไม่รับรอง Historical PIT และไม่รับรองผลเชิงสาเหตุ</p></div>'))
    assert len(pages)==10
    css='''@page{size:A4;margin:0}*{box-sizing:border-box}body{margin:0;background:#e8eef0;color:#132d38;font:15px/1.65 "Leelawadee UI",Tahoma,sans-serif}.page{width:210mm;min-height:297mm;padding:15mm 17mm 18mm;background:white;margin:8mm auto;position:relative;break-after:page}.page:last-child{break-after:auto}header{font-size:12px;border-bottom:1px solid #b8cbd1;padding-bottom:8px}header span,footer span{float:right}h1{font-size:25px;line-height:1.35;margin:22px 0 18px}h2{font-size:18px;margin:16px 0 9px}p{margin:10px 0}.lead{font-size:20px}.notice{padding:12px 16px;background:#edf5f6;border-left:4px solid #087f8c;margin:16px 0}li{margin:11px 0}svg{display:block;width:100%;height:auto;font:16px "Leelawadee UI",Tahoma,sans-serif}svg text{fill:#132d38}.scatter svg{height:185px}.small{font-size:12px}table{border-collapse:collapse;width:100%;font-size:13px}td,th{padding:7px 8px;border-bottom:1px solid #dbe3e9;text-align:left}th{background:#edf3f5}footer{position:absolute;bottom:10mm;left:17mm;right:17mm;font-size:11px;color:#49616c;border-top:1px solid #dbe3e9;padding-top:6px}@media print{body{background:white}.page{margin:0;height:297mm;min-height:0}}'''
    textfile(out/'final_advisor_report_th.html','<!doctype html><html lang="th"><meta charset="utf-8"><title>รายงานทุเรียนจันทบุรีสำหรับอาจารย์</title><style>'+css+'</style>'+''.join(pages)+'</html>')
    dictionary='''# Data Dictionary

หน่วยหลัก analysis_results.csv คือจังหวัดจันทบุรี–ปี ไม่ใช่สวนหรือ Polygon ช่อง CSV ว่าง = null ไม่ใช่ศูนย์

|Field|Unit / definition|Source / calculation|
|---|---|---|
|province_code / year_ce / year_be|รหัสจังหวัด / ค.ศ. / พ.ศ.|22 / ปีต้นทาง / CE+543|
|yield_kg_rai|กก./ไร่พื้นที่ให้ผล|OAE reported ไม่คำนวณเติม|
|temperature_c / humidity_pct|°C / % RH2M|POWER ค่าเดือนถ่วงวัน ม.ค.–ธ.ค. ปีเดียวกับ Target; descriptive only|
|rain_mm|มม./ปี|Σ PRECTOTCORR mm/day × days; descriptive only|
|bearing_area_rai / production_tonnes|ไร่ / ตัน|OAE reported|
|issue_date|YYYY-03-01|กฎจำลองใหม่ ไม่อ้างว่าใช้ในอดีตจริง|
|feature_observation_end|YYYY-1-12-31|Observation cutoff ไม่ใช่ Publication evidence|
|yield_lag2_kg_rai|กก./ไร่|Yield Y−2; ว่างเมื่อไม่มีต้นทาง|
|t2m_lag1_c / rain_lag1_mm / rh_lag1_pct|°C / mm / %|อากาศ Y−1 ม.ค.–ธ.ค.; Forecast inputs|
|split|historical / inner / test|Protocol ที่ล็อกไว้|
|availability_status|unverified_publication_vintage|ไม่รับรอง Historical PIT|
|weather_footprint_year_be|2568|รุ่น A403 ใช้ spatial weights ย้อนหลัง ไม่ใช่สวนคงที่จริง|

model_evaluation.csv หนึ่งแถวต่อ split–model–year รวม 30 แถว (Inner 4×3 และ Test 6×3)

|Field|Definition|
|---|---|
|record_type / split / model|prediction / inner หรือ test / lag2_baseline, training_mean, ridge|
|alpha|Ridge penalty; null สำหรับ Baselines|
|n_train / fit_label_year_max|จำนวน Training rows / ปีล่าสุด ≤ Target−2|
|actual_kg_rai / predicted_kg_rai|กก./ไร่; ไม่ Clamp|
|error_kg_rai / absolute_error_kg_rai|Prediction−Actual / absolute value|
|issue_date / feature_observation_end / availability_status|นิยามเดียวกับตารางหลัก|

MAE=mean(abs(error)), RMSE=sqrt(mean(error²)), Bias=mean(error). Inner grid evaluation แยกเพื่อไม่ปน candidate กับ selected scores.
tha_mai_soil_summary.csv หนึ่งแถวต่อหน่วยแผนที่ดิน; area_rai และ denominator_rai ใช้ Accepted Proposal เดิม share_pct=100×area/denominator รวม SC/W ใน covered และ Uncovered แยกในรายงาน ไม่รวม Mixed-use.
district_case_study.csv หนึ่งแถวต่ออำเภอ–ปี reported_yield_kg_rai จาก PDF หน้า5 Table2; status=source_transcription_only_finality_unverified ไม่ยกค่าเฉลี่ยไปให้สมาชิก/สวน.
'''
    textfile(out/'data_dictionary.md',dictionary)
    src='# Source Manifest — Internal Only / Public HOLD\n\nPublication Date ไม่ยืนยัน; Retrieval/Data Date ไม่ใช้แทน Publication Date. Land-use2568/Soil2561 ไม่เท่ากับ observation date ของสวนย้อนหลัง.\n\n'+''.join(f'- `{k}`\n  SHA256 `{v}`\n' for k,v in inputhash.items())+'\n## Lineage และสิทธิ์\n\nPhase9D Source Manifest ผูก OAE/POWER/raw LDD; Phase11B ผูก PDF SHA และ Page/Table/Row anchors. Source cells OAE อยู่ analysis_summary.json. Licenses ให้ตรวจรายการต้นทาง ไม่ตีความ Government-hosted ว่าเป็น CC0. Derived LDD/Geometry public release HOLD, OAE/POWER vintages unverified. ไม่มีการดาวน์โหลดใหม่.\n'
    textfile(out/'source_manifest.md',src)
    textfile(out/'README_TH.md',f'''# ผลวิเคราะห์ทุเรียนจันทบุรีสำหรับ Selected Topics in AI

เปิด final_advisor_report_th.pdf หรือ HTML (10 หน้า) ตัวเลขทุกกราฟใช้ข้อมูลเดิมในเครื่อง ไม่แก้ Source/Production/API/Frontend

## ตอบอะไรได้
แนวโน้มจังหวัด2006–2025, ความสัมพันธ์รายปีแบบพรรณนา, ประสิทธิภาพ Ridge เทียบ Lag2 และ Training Mean ในการจำลองตามเวลา, กรณีรายอำเภอสองปี และองค์ประกอบดินท่าใหม่

## ยังตอบไม่ได้
เหตุและผลจากอุณหภูมิ, Yield รายสวน, ดินที่ทำให้ Yield สูงข้ามอำเภอ, Forecast ที่รับรอง PIT, Confidence Interval ที่สอบเทียบจริง และสถานการณ์โลกร้อนเชิงสาเหตุ. สถานะ Scenario: {scenario_status}

## การทำซ้ำ
ใช้ Python ที่มี numpy และ openpyxl จาก C:/project:
`python -m tools.build_final_advisor_analysis --output-name final_reproduction_new_id`
โปรแกรมปฏิเสธโฟลเดอร์ที่มีอยู่แล้ว ก่อนรันตรวจ Source Seals/Hashes และบันทึก protocol.json ก่อน Fit ไม่เลือก Alpha จาก Test. reproduce.py เป็นสำเนา Source Script ใช้ helper เดิมที่บันทึก Hash. HTML เกิดซ้ำจาก Script; PDF ใช้ render_pdf.cjs ผ่าน Edge/Playwright ที่มีในเครื่อง (ไม่ติดตั้งหรือดาวน์โหลด). ปี Test เคยใช้พัฒนาแล้ว ไม่เป็นอิสระ.

## ไฟล์
analysis_results.csv: 20 จังหวัด–ปี; model_evaluation.csv: 30 Predictions; inner_grid_evaluation.csv: 12 Candidate Validation Predictions; analysis_summary.json: Summary/Gate/QC; ridge_fold_parameters.json: Scalers/Coefficients; soil/district CSV แยกหน่วยข้อมูล; source_manifest.md และ input_manifest.json: Paths/Hashes; qc_report.json: Checks/Protected Verification. CSV UTF-8 BOM ช่องว่างเป็น null.

## วิธีอ่าน
กราฟพรรณนาใช้ Weather ปีเดียวกับ Yield ไม่เข้าโมเดล. โมเดลใช้ Weather Y−1, Yield Y−2; fit Labels≤Y−2. วัน1มีนาคมเป็นกฎจำลองใหม่ ไม่เปลี่ยน Protocol/คะแนนของ Phaseเดิม. ข้อมูล OAE/POWER revised snapshots และ static footprint ยังไม่ PIT. ไม่ใช้ความสัมพันธ์เป็นคำแนะนำทางเกษตร. Public HOLD เสมอ.

## สภาพแวดล้อม
Python {platform.python_version()}, NumPy {np.__version__}; closed-form Ridge ใช้ train-only scaling และ unpenalized intercept ไม่ติดตั้ง sklearn. งานตรวจสอบเฉพาะขั้นคำนวณนี้ ไม่รัน Repository Tests ที่ไม่เกี่ยวข้อง. ใช้หลัก Ponytail แยก Script เล็กจาก Production และสกิล Spreadsheets/PDF เพื่อ Units/Provenance/Visual QA.
''')
    textfile(out/'reproduce.py',Path(__file__).read_text(encoding='utf-8'))
    assert all(sha(ROOT/k)==v for k,v in protected.items()),'Protected changed'
    assert all(sha(ROOT/k)==v for k,v in inputhash.items()),'Source changed'
    save(out/'qc_report.json',dict(status='pass_with_limitations',checks=checks,protected_files_checked=len(protected),protected_hash_mismatches=[],input_files_checked=len(inputhash),production_before=production,production_after=sha(MODEL),publication_vintage='unverified',public_release='HOLD',scenario_status=scenario_status))
    print(json.dumps(dict(output=str(out),metrics=summary,selected_alpha=selected,gate=gate,checks=len(checks)),ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--output-name',default='final_20261012_v1'); args=p.parse_args()
    assert Path(args.output_name).name==args.output_name and args.output_name not in ('.','..')
    build(ROOT/'research_data/internal_research/final_advisor'/args.output_name)
