"""Phase 11B offline, visually verified transcription; no GIS join/network/ML."""
from collections import defaultdict
from html import escape
import math
from pathlib import Path
from openpyxl import load_workbook
import shapefile
from tools import build_tha_mai_spatial_pilot as p

OUT=Path("research_data/internal_research/chanthaburi_district_yield/phase11b_20261012_v2")
PDF_SHA="84d4cc284f65f77ccf5c4a42ae7df6ab21f730cee6913b9fdc13ce8af21394ab"
FIELDS=["standing_area_rai","bearing_area_rai","production_tonnes","reported_yield_kg_rai"]
UNITS=dict(zip(FIELDS,["rai","rai","tonne","kg/rai"]))
COLUMNS=dict(zip(FIELDS,["เนื้อที่ยืนต้น (ไร่)","เนื้อที่ให้ผล (ไร่)","ผลผลิต (ตัน)","ผลผลิตต่อเนื้อที่ให้ผล (ก.ก./ไร่)"]))
# Two years, each in standing / bearing / production / reported-yield order.
# Transcribed from rendered PDF page 5, Table 2, not inferred from OAE.
TRANSCRIPTION=[
 ("เมืองจันทบุรี",[10615,8694,16736,1925],[12553,8969,13193,1471]),
 ("ขลุง",[80294,59135,103900,1757],[83316,64217,103325,1609]),
 ("ท่าใหม่",[105816,77130,171229,2220],[109116,80690,130314,1615]),
 ("โป่งน้ำร้อน",[36572,16457,23073,1402],[37021,19337,20420,1056]),
 ("มะขาม",[40387,27508,54136,1968],[48582,28540,46006,1612]),
 ("แหลมสิงห์",[1819,1586,2737,1726],[1896,1616,2490,1541]),
 ("สอยดาว",[5522,1061,960,905],[6161,989,716,724]),
 ("แก่งหางแมว",[42793,17676,32701,1850],[48417,22464,28754,1280]),
 ("นายายอาม",[34265,24957,45621,1828],[35550,28423,38428,1352]),
 ("เขาคิชฌกูฏ",[54963,39444,87368,2215],[56940,40028,65046,1625])]
PROVINCE={2566:dict(zip(FIELDS,[413046,273648,538461,1968])),
          2567:dict(zip(FIELDS,[439552,295273,448692,1520]))}
FLAGS=["transcription_verified_against_rendered_page","year_basis_unverified","finality_unverified",
       "publication_date_unverified","secondary_table_citing_oae6","license_pending","public_release_hold",
       "admin_effective_date_unverified","printed_integer_precision"]


def match_name(name,admin):
    matches=[r for r in admin if r["AMPHOE_T"].removeprefix("อ.").strip()==name.strip()]
    codes={str(r["AMPHOE_IDN"]) for r in matches}
    if len(codes)!=1: return None,"ambiguous" if codes else "unmatched"
    return codes.pop(),"verified_normalized_exact"


def yield_check(row):
    a,prod,y=(row[f] for f in ["bearing_area_rai","production_tonnes","reported_yield_kg_rai"])
    if any(v is None for v in [a,prod,y]) or a<=0:
        return {"status":"unavailable","derived_yield_kg_rai":None,"reason":"missing_or_nonpositive_bearing"}
    ratio=prod*1000/a
    # Allow only rounding intervals supported by the printed whole-unit columns.
    lower=max(0,prod-.5)*1000/(a+.5)
    upper=(prod+.5)*1000/(a-.5) if a>.5 else float("inf")
    return {"status":"rounding_consistent" if y+.5>=lower and y-.5<=upper else "inconsistent",
            "derived_yield_kg_rai":ratio,"reported_minus_derived_kg_rai":y-ratio,
            "rounding_interval_kg_rai":[lower,upper]}


def validate(rows):
    keys=[(r["district_code"],r["year_be"]) for r in rows]
    if len(set(keys))!=len(keys): raise ValueError("Duplicate district-year")
    if set(keys)!={(str(code),y) for code in range(2201,2211) for y in [2566,2567]}:
        raise ValueError("Missing/extra district-year coverage")
    for row in rows:
        if row["crosswalk_status"]!="verified_normalized_exact": raise ValueError("Unverified crosswalk")
        if row["year_ce"]!=row["year_be"]-543: raise ValueError("Year conversion")
        for field in FIELDS:
            value=row[field]
            if row["units"].get(field)!=UNITS[field]: raise ValueError("Unit mismatch")
            if value is not None and (not math.isfinite(value) or value<0): raise ValueError("Negative/nonfinite value")
            if row["missingness"][field]!=("present" if value is not None else "missing_in_source"):
                raise ValueError("Missingness mismatch")
        if not row["source_anchor"].get("district_row_name"): raise ValueError("Missing source anchor")
        if yield_check(row)["status"]=="inconsistent": raise ValueError("Yield not rounding-consistent")
    return True


def difference(value,reference):
    return {"difference_signed":value-reference,"difference_absolute":abs(value-reference),
            "difference_pct_of_reference":100*(value-reference)/reference if reference else None,
            "absolute_difference_pct_of_reference":100*abs(value-reference)/reference if reference else None}


def extract():
    root=p.ROOT; out=root/OUT
    assert not (out/"district_year.json").exists(),"Refuse overwrite"
    assert p.sha(out/"source.pdf")==PDF_SHA
    admin_path=root/p.ADMIN.with_suffix(".dbf")
    admin=[]
    for index,record in enumerate(shapefile.Reader(dbf=str(admin_path),encoding="utf-8").iterRecords()):
        admin.append({**record.as_dict(),"source_record_0based":index})
    rows=[]; crosswalk=[]; lineage=[]
    for order,(name,first,second) in enumerate(TRANSCRIPTION,1):
        code,status=match_name(name,admin)
        assert status=="verified_normalized_exact",name
        matched=[r for r in admin if str(r["AMPHOE_IDN"])==code]
        crosswalk.append({"table_district_name":name,"district_code":code,"status":status,
            "normalization":"strip leading อ. from local admin name only",
            "admin_names":sorted({r["AMPHOE_T"] for r in matched}),
            "admin_source_records_0based":[r["source_record_0based"] for r in matched],
            "admin_effective_date":None,"admin_sha256":p.sha(admin_path),"admin_path":str(p.ADMIN.with_suffix('.dbf'))})
        for year,values in [(2566,first),(2567,second)]:
            anchor={"pdf_sha256":PDF_SHA,"pdf_page_1based":5,"printed_page_number":None,
                    "table_number":2,"district_row_ordinal_excluding_province":order,"district_row_name":name,
                    "year_column_label_be":str(year)}
            row={"province_code":"22","province_name":"จันทบุรี","district_code":code,"district_name":name,
                "year_be":year,"year_ce":year-543,"year_basis":"unverified_year_label_only",
                **dict(zip(FIELDS,values)),"units":UNITS.copy(),"source_anchor":anchor,
                "missingness":{field:"present" if v is not None else "missing_in_source" for field,v in zip(FIELDS,values)},
                "quality_flags":FLAGS.copy(),"crosswalk_status":status,"finality":"unverified",
                "publication_date":None,"acceptance_scope":"source_transcription_only","training_eligible":False}
            rows.append(row)
            for field,value in zip(FIELDS,values):
                lineage.append({"district_code":code,"year_be":year,"field":field,"value":value,
                    "source_column":COLUMNS[field],"printed_value":f"{value:,}" if value is not None else None,
                    "method":"manual_visual_transcription_from_original_render","anchor":anchor})
            for field,origin in [("district_name","PDF table row label"),("year_be","PDF column label"),
                                 ("year_ce","year_be minus 543; not confirmation of calendar basis"),
                                 ("district_code","verified local DBF name crosswalk; code not printed in PDF")]:
                lineage.append({"district_code":code,"year_be":year,"field":field,"method":origin,
                                "anchor":anchor,"crosswalk_reference":code if field=="district_code" else None})
    validate(rows)
    workbook=root/"research_data/external/oae_2026-09-17/durian_province.xlsx"
    assert p.sha(workbook)=="6a3ae44026b5ef8037a8bb091852281d890168a8d795c9c44d5cde9a48613d4c"
    w=load_workbook(workbook,read_only=True,data_only=True); sheet=w["durian"]
    names=dict(zip(["เนื้อที่ยืนต้น","เนื้อที่ให้ผล","ผลผลิต","ผลผลิตต่อเนื้อที่ให้ผล"],FIELDS))
    oae={}; oae_anchors=[]
    for index,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
        year,crop,province,commodity,subcom,item,value,unit,data_date=values
        if str(province).strip()=="จันทบุรี" and int(year) in [2566,2567] and item in names:
            field=names[item]; key=(int(year),field); assert key not in oae
            assert crop==year and commodity=="ทุเรียน"
            assert unit=={"rai":"ไร่","tonne":"ตัน","kg/rai":"กิโลกรัมต่อไร่"}[UNITS[field]]
            oae[key]=value; oae_anchors.append({"year_be":year,"field":field,"value":value,
                "sheet":"durian","row_1based":index,"value_cell":f"G{index}","data_date":str(data_date),
                "publication_date":None,"vintage_status":"local_reference_vintage_history_unverified"})
    w.close(); assert len(oae)==8
    reconciliation=[]
    for year in [2566,2567]:
        for field in FIELDS:
            total=sum(r[field] for r in rows if r["year_be"]==year) if field!="reported_yield_kg_rai" else None
            reconciliation.append({"year_be":year,"year_ce":year-543,"field":field,"unit":UNITS[field],
                "district_sum":total,"pdf_province_reported":PROVINCE[year][field],"oae_reference":oae[(year,field)],
                "district_sum_minus_pdf_province":total-PROVINCE[year][field] if total is not None else None,
                "aggregation":"sum" if total is not None else "ratio_of_production_sum_to_bearing_sum; never sum yields",
                "derived_province_yield":PROVINCE[year]['production_tonnes']*1000/PROVINCE[year]['bearing_area_rai'] if total is None else None,
                **difference(PROVINCE[year][field],oae[(year,field)])})
    p.write(out/"district_year.json",{"schema_version":1,"rows":rows,"public_release":"HOLD"})
    p.write(out/"province_rows.json",[{"year_be":year,**record,"anchor":{"pdf_page_1based":5,"table_number":2,"row_name":"จันทบุรี"}} for year,record in PROVINCE.items()])
    p.write(out/"district_crosswalk.json",crosswalk); p.write(out/"field_lineage.json",lineage)
    p.write(out/"oae_reference_anchors.json",oae_anchors); p.write(out/"reconciliation.json",reconciliation)
    checks=[{"district_code":r['district_code'],"year_be":r['year_be'],**yield_check(r)} for r in rows]
    p.write(out/"quality_report.json",{"status":"accepted_transcription_only","rows":len(rows),"districts":10,
        "years_be":[2566,2567],"numeric_cells":80,"missing_numeric_cells":sum(r[f] is None for r in rows for f in FIELDS),
        "yield_checks":checks,"final_actual_verified":False,"year_basis_verified":False,
        "publication_date_verified":False,"license_status":"unverified","public_release":"HOLD"})
    manifest=p.read_json(out/"download_verification.json")
    manifest.update({"source_role":"secondary_research_document_hosted_on_DOAE_domain",
        "document_title":"แนวทางการประยุกต์ใช้ข้อมูลทะเบียนเกษตรกรในการประเมินสถานการณ์ไม้ผลของเจ้าหน้าที่สำนักงานเกษตรอำเภอในพื้นที่จังหวัดจันทบุรี",
        "author":"นายศฤงคาร รักร่วม","author_affiliation":"สำนักงานเกษตรจังหวัดจันทบุรี",
        "table_attribution":"สำนักงานเศรษฐกิจการเกษตรที่ 6, 2567","table_page_1based":5,"table_number":2,
        "table_title":"ตารางทุเรียน เนื้อที่ยืนต้น เนื้อที่ให้ผล ผลผลิต ผลผลิตต่อไร่ รายอำเภอ ปี 2566 - 2567",
        "table_image_header":"รายอำเภอ ภาคตะวันออก ปี 2566 - 2567","document_section":"ข้อมูลเอกภาพไม้ผลที่สำคัญของจังหวัดจันทบุรี",
        "year_basis":"unverified","finality":"unverified","license_status":"unverified_public_hold",
        "publication_date_evidence":"No explicit release date for this table. URL month, HTTP Last-Modified and PDF metadata are not publication proof.",
        "study_period_context":"page 3: July-August 2567; not asserted as table valid period or publication date",
        "oae_workbook_sha256":p.sha(workbook),"admin_dbf_sha256":p.sha(admin_path),
        "page_image_sha256":p.sha(out/"source_page-05.png"),"code_sha256":p.sha(__file__)})
    p.write(out/"source_manifest.json",manifest)
    p.write(out/"data_dictionary.json",{"primary_key":["district_code","year_be"],"fields":{
        **{field:{"unit":UNITS[field],"definition":COLUMNS[field],"source":"PDF page 5 Table 2","missing":"null, not zero"} for field in FIELDS},
        "district_code":{"source":"local admin DBF crosswalk","meaning":"code verified against local layer; historical effective date unverified"},
        "year_be":{"meaning":"printed BE year label; crop/harvest/calendar basis not specified"},
        "year_ce":{"calculation":"year_be minus 543; no calendar reinterpretation"},
        "derived_yield_kg_rai":{"unit":"kg/rai","calculation":"production_tonnes*1000/bearing_area_rai; QC only, does not overwrite reported yield"},
        "missingness":{"values":["present","missing_in_source"]},"source_anchor":{"meaning":"PDF hash, physical page, table, named district row, ordinal, BE year column"},
        "quality_flags":{"meaning":"transcription and temporal/finality/license limits"}},
        "standing_area_warning":"standing/planted area is not bearing area; no GIS area is used as denominator"})
    table="".join(f"<tr><td>{r['district_code']}</td><td>{escape(r['district_name'])}</td><td>{r['year_be']} / {r['year_ce']}</td>"+"".join(f"<td>{r[field]:,}</td>" if r[field] is not None else "<td>missing</td>" for field in FIELDS)+"<td>5 / 2 / "+str(r['source_anchor']['district_row_ordinal_excluding_province'])+"</td></tr>" for r in rows)
    html=f"<!doctype html><html lang='th'><meta charset='utf-8'><title>Phase 11B district yield preview</title><style>body{{font-family:Tahoma,Arial;color:#173742;margin:30px}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{padding:8px;border-bottom:1px solid #ccd7db}}td{{text-align:right}}td:nth-child(2){{text-align:left}}th{{background:#e8f0f2}}p{{line-height:1.7}}</style><h1>จันทบุรี: ทุเรียนรายอำเภอ 2566–2567</h1><p>Internal Only / Public HOLD · รับรองการถอดตารางเท่านั้น ไม่รับรอง Final Actual หรือปีเพาะปลูก/เก็บเกี่ยว/ปฏิทิน · Source: DOAE-hosted PDF, page 5, Table 2; อ้าง สศท.6, 2567</p><table><tr><th>รหัส</th><th>อำเภอ</th><th>พ.ศ. / ค.ศ.</th><th>ยืนต้น (ไร่)</th><th>ให้ผล (ไร่)</th><th>ผลผลิต (ตัน)</th><th>Yield (กก./ไร่)</th><th>หน้า/ตาราง/แถว</th></tr>{table}</table><p>รหัสตรวจจาก Local Admin DBF ไม่ได้พิมพ์อยู่ในตารางต้นฉบับ ไม่มี Join GIS/Weather, ML หรือปรับตัวเลขตาม OAE</p></html>"
    (out/"dataset_preview.html").write_text(html,encoding="utf-8")
    print({"rows":len(rows),"missing_numeric":0,"yield_statuses":sorted({c['status'] for c in checks}),
           "nonzero_oae_differences":[r for r in reconciliation if r['difference_absolute']]})


if __name__=="__main__": extract()
