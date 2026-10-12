"""Offline descriptive report only. No model fitting, GIS operations, or source writes."""
import csv
import json
import math
import hashlib
from html import escape as e
from pathlib import Path
from statistics import mean
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'research_data/internal_research/chanthaburi_eda/eda_20261012_v1'
P9=ROOT/'research_data/internal_research/four_provinces/phase9d_20261012_v1'
P10=ROOT/'research_data/internal_research/tha_mai_spatial/phase10f_20261012_internal_accepted_proposal_v1'
P11=ROOT/'research_data/internal_research/chanthaburi_district_yield/phase11b_20261012_v2'
MODEL=ROOT/'research_data/five_province_history/training/ridge_vpd_final_model.json'


def sha(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()


def read(path):return json.loads(path.read_text(encoding='utf-8'))
def rows(path):return list(csv.DictReader(path.open(encoding='utf-8-sig',newline='')))
def save(path,value):
    with path.open('x',encoding='utf-8') as f:json.dump(value,f,ensure_ascii=False,indent=2)
def seal_check(folder):
    m=read(folder/'output_manifest.json')
    assert all(sha(folder/k)==v for k,v in m['files'].items()),str(folder)
    return m


def corr(x,y):
    assert len(x)==len(y)>1 and all(math.isfinite(v) for v in x+y)
    mx,my=mean(x),mean(y)
    return sum((a-mx)*(b-my) for a,b in zip(x,y))/math.sqrt(sum((a-mx)**2 for a in x)*sum((b-my)**2 for b in y))


def svg_plot(xs,ys,title,xlabel,ylabel,xlim,ylim,color,line=False,years=None):
    W,H=940,340; L,R,T,B=95,30,35,65
    X=lambda x:L+(x-xlim[0])/(xlim[1]-xlim[0])*(W-L-R)
    Y=lambda y:H-B-(y-ylim[0])/(ylim[1]-ylim[0])*(H-T-B)
    s=[f'<svg role="img" aria-label="{e(title)}" viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg"><title>{e(title)}</title>']
    for k in range(5):
        v=ylim[0]+(ylim[1]-ylim[0])*k/4;py=Y(v)
        s.append(f'<path d="M{L},{py}H{W-R}" stroke="#dbe3e9"/><text x="{L-12}" y="{py+5}" text-anchor="end">{v:,.1f}</text>')
    ticks=sorted(set([2006,2010,2014,2018,2022,2025])) if line else [xlim[0]+(xlim[1]-xlim[0])*k/4 for k in range(5)]
    for v in ticks:
        s.append(f'<text x="{X(v)}" y="{H-B+28}" text-anchor="middle">{int(v) if line else f"{v:,.1f}"}</text>')
    s.append(f'<text x="{L}" y="21">{e(ylabel)}</text><text x="{(W+L-R)/2}" y="{H-8}" text-anchor="middle">{e(xlabel)}</text>')
    if line:s.append(f'<polyline fill="none" stroke="{color}" stroke-width="3" points="'+ ' '.join(f'{X(x):.2f},{Y(y):.2f}' for x,y in zip(xs,ys))+'"/>')
    for i,(x,y) in enumerate(zip(xs,ys)):
        year=int(xs[i]) if line else years[i]
        s.append(f'<circle tabindex="0" cx="{X(x):.2f}" cy="{Y(y):.2f}" r="5" fill="{color}" stroke="white"><title>ปี {year}: {xlabel} {x:,.2f}; {ylabel} {y:,.2f}</title></circle>')
        if not line and year in (2016,2022,2024):
            s.append(f'<text x="{X(x)+9}" y="{Y(y)-8}">{year}</text>')
    return ''.join(s)+'</svg>'


def table(headers,data):
    return '<div class="table-scroll"><table><thead><tr>'+''.join(f'<th>{e(h)}</th>' for h in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join(f'<td>{e(str(v))}</td>' for v in r)+'</tr>' for r in data)+'</tbody></table></div>'


def build():
    assert not OUT.exists(),'Refuse overwrite'
    before=read(ROOT/'research_data/internal_research/chanthaburi_district_yield/phase11e_20261012_v1/preflight_protected_hashes.json')
    assert all(sha(ROOT/k)==v for k,v in before.items()),'Protected preflight mismatch'
    for p in (P9,P10,P11):seal_check(p)
    assert sha(MODEL)=='b4b9e0985e6ba5e6ee371d8988cad5e65dee031b345eb0cf13a8695f91658680'
    annual=sorted([r for r in rows(P9/'province_year.csv') if r['province_code']=='22'],key=lambda r:int(r['year_ce']))
    assert len(annual)==20 and [int(r['year_ce']) for r in annual]==list(range(2006,2026))
    monthly=[r for r in rows(P9/'weather_profile_month.csv') if r['province_code']=='22']
    assert len(monthly)==240 and len({(r['year_ce'],r['month']) for r in monthly})==240
    for r in annual:
        mm=[m for m in monthly if m['year_ce']==r['year_ce']]
        assert sorted(int(m['month']) for m in mm)==list(range(1,13))
        days=sum(int(m['days']) for m in mm)
        assert abs(sum(float(m['T2M'])*int(m['days']) for m in mm)/days-float(r['weather_t2m_mean_c']))<1e-6
        assert abs(sum(float(m['PRECTOTCORR'])*int(m['days']) for m in mm)-float(r['weather_precip_total_mm']))<1e-6
        assert all(float(m['T2M_area_coverage'])>.999999 and float(m['PRECTOTCORR_area_coverage'])>.999999 for m in mm)
    workbook=ROOT/'research_data/external/oae_2026-09-17/durian_province.xlsx'
    w=load_workbook(workbook,read_only=True,data_only=True); source=[]
    mapping={'ผลผลิตต่อเนื้อที่ให้ผล':'yield_kg_per_rai','เนื้อที่ให้ผล':'bearing_area_rai','เนื้อที่ยืนต้น':'planted_area_rai','ผลผลิต':'production_tonnes'}
    for index,r in enumerate(w['durian'].iter_rows(min_row=2,values_only=True),2):
        year,crop,province,commodity,subcom,item,v,unit,date=r
        if province=='จันทบุรี' and commodity=='ทุเรียน' and item in mapping and 2549<=int(year)<=2568:
            match=[x for x in annual if int(x['year_ce'])==int(year)-543];assert len(match)==1
            assert float(match[0][mapping[item]])==float(v)
            assert unit=={'yield_kg_per_rai':'กิโลกรัมต่อไร่','production_tonnes':'ตัน','bearing_area_rai':'ไร่','planted_area_rai':'ไร่'}[mapping[item]]
            source.append(dict(year_be=year,field=mapping[item],cell=f'G{index}',sheet='durian',value=v,unit=unit,data_date=str(date)))
    w.close();assert len(source)==80
    district=read(P11/'district_year.json')['rows'];assert len(district)==20
    assert {(r['district_code'],r['year_be']) for r in district}=={(str(c),y) for c in range(2201,2211) for y in (2566,2567)}
    data=read(P10/'dataset.json')['tables'];soil=data['soil_composition'];assert len(soil)==43
    den=soil[0]['denominator_pure_rai'];unc=sum(r['area_rai'] for r in data['uncovered'])
    assert all(r['district_id']=='2203' and r['denominator_pure_rai']==den for r in soil)
    assert abs(sum(r['area_rai'] for r in soil)+unc-den)<=.1
    assert abs(sum(r['share_of_all_pure_pct'] for r in soil)+100*unc/den-100)<100*.1/den
    assert all(abs(r['share_of_all_pure_pct']-100*r['area_rai']/den)<1e-7 for r in soil)
    names={r['soil_map_code']:r['soil_name_source'] for r in rows(P9/'soil_composition.csv') if r['province_code']=='22'}
    yy=[int(r['year_ce']) for r in annual];y=[float(r['yield_kg_per_rai']) for r in annual]
    t=[float(r['weather_t2m_mean_c']) for r in annual];rain=[float(r['weather_precip_total_mm']) for r in annual]
    rt,rr=corr(t,y),corr(rain,y)
    low,high=min(annual,key=lambda r:float(r['yield_kg_per_rai'])),max(annual,key=lambda r:float(r['yield_kg_per_rai']))
    a23,a24=[r for r in annual if int(r['year_ce'])==2023][0],[r for r in annual if int(r['year_ce'])==2024][0]
    delta=(float(a24['yield_kg_per_rai'])/float(a23['yield_kg_per_rai'])-1)*100
    decrease=sum(next(r['reported_yield_kg_rai'] for r in district if r['district_code']==str(c) and r['year_be']==2567)<next(r['reported_yield_kg_rai'] for r in district if r['district_code']==str(c) and r['year_be']==2566) for c in range(2201,2211))
    top=sorted(soil,key=lambda r:-r['area_rai']);top3=sum(r['share_of_all_pure_pct'] for r in top[:3])
    findings=[f"Yield จังหวัดสูงสุดในปี {int(high['year_ce'])+543} ({high['year_ce']}) {float(high['yield_kg_per_rai']):,.0f} กก./ไร่ และต่ำสุดปี {int(low['year_ce'])+543} ({low['year_ce']}) {float(low['yield_kg_per_rai']):,.0f} กก./ไร่ แนวโน้มมีขึ้นลง ไม่เพิ่มสม่ำเสมอทุกปี",
        f"ปี 2567 (2024) Yield จังหวัดลด {abs(delta):.1f}% จากปี 2566 ขณะที่อุณหภูมิเฉลี่ยทั้งปี {float(a24['weather_t2m_mean_c']):.2f} °C และฝนรวม {float(a24['weather_precip_total_mm']):,.0f} มม. การเกิดพร้อมกันไม่พิสูจน์ว่าอากาศทำให้ Yield ลด",
        f"Pearson r ระหว่าง Yield กับอุณหภูมิเฉลี่ยรายปี = {rt:.3f} และกับฝนรวมรายปี = {rr:.3f} จาก 20 จังหวัด–ปี ค่าดิบนี้ไม่ควบคุมแนวโน้มเวลา พื้นที่ให้ผล พันธุ์ อายุสวน หรือการจัดการ และไม่ใช่ผลทดสอบเหตุและผล",
        f"ในตารางรายอำเภอ Yield ปี 2567 ต่ำกว่าปี 2566 ทั้ง {decrease}/10 อำเภอ เป็นข้อสังเกตที่ควรถามต้นทางเรื่องฤดู/รุ่นข้อมูลก่อนอธิบายทางชีวภาพ ตารางยังไม่รับรอง Final Actual",
        f"Pure A403 ท่าใหม่มี {len(soil)} รหัสหน่วยแผนที่ดิน รหัสใหญ่สุด {top[0]['soil_unit']} มี {top[0]['share_of_all_pure_pct']:.2f}% และสามรหัสแรก {top3:.2f}% สะท้อนความหลากหลายของพื้นที่บนแผนที่ ไม่บอกว่าดินรหัสใดทำให้ Yield สูง และไม่เป็นตัวแทนทุกอำเภอ"]
    rank_data=[]
    for year in (2566,2567):
        ranked=sorted([r for r in district if r['year_be']==year],key=lambda r:-r['reported_yield_kg_rai'])
        for rank,r in enumerate(ranked,1):rank_data.append([year,rank,r['district_name'],f"{r['reported_yield_kg_rai']:,}",f"{r['production_tonnes']:,}",f"{r['bearing_area_rai']:,}"])
    trend=''.join('<h3>'+label+'</h3>'+svg_plot(yy,vals,label,'ปี ค.ศ.',unit,(2006,2025),lim,c,line=True) for label,vals,unit,lim,c in [
        ('ผลผลิตต่อพื้นที่ให้ผล',y,'กก./ไร่',(0,2500),'#2563eb'),('อุณหภูมิอากาศเฉลี่ย 2 เมตร',t,'°C · แกนไม่เริ่มจากศูนย์',(25,28),'#087f8c'),('ฝนรวมรายปี',rain,'มม./ปี',(0,3500),'#8b5e34')])
    scatter=svg_plot(t,y,'อุณหภูมิรายปีกับ Yield','อุณหภูมิเฉลี่ย ม.ค.–ธ.ค. (°C)','Yield (กก./ไร่)',(25.5,27.5),(0,2500),'#087f8c',years=yy)+svg_plot(rain,y,'ฝนรายปีกับ Yield','ฝนรวม ม.ค.–ธ.ค. (มม.)','Yield (กก./ไร่)',(1800,3200),(0,2500),'#2563eb',years=yy)
    # Other is a display group only; SC and W remain explicit and are already inside Covered.
    bars=[(r['soil_unit'],r['share_of_all_pure_pct']) for r in top if r['soil_unit'] not in ('SC','W')][:8]
    included={k for k,v in bars};bars += [('Other หน่วยอื่น',sum(r['share_of_all_pure_pct'] for r in soil if r['soil_unit'] not in included|{'SC','W'}))]
    bars += [(k,next((r['share_of_all_pure_pct'] for r in soil if r['soil_unit']==k),0)) for k in ('SC','W')]+[('Uncovered',100*unc/den)]
    maxbar=max(v for k,v in bars);bar=['<svg viewBox="0 0 940 535" role="img" aria-label="สัดส่วนหน่วยแผนที่ดินท่าใหม่"><title>สัดส่วนหน่วยแผนที่ดินท่าใหม่</title>']
    for i,(name,v) in enumerate(bars):
        py=23+i*40;bar.append(f'<text x="180" y="{py+19}" text-anchor="end">{e(name)}</text><rect x="195" y="{py}" width="{v/maxbar*590:.4f}" height="25" fill="#087f8c"/><text x="{205+v/maxbar*590:.4f}" y="{py+19}">{v:.3f}%</text>')
    bar.append('</svg>')
    soiltable=table(['รหัสหน่วยแผนที่','ชื่อจาก Source เมื่อมี','พื้นที่ (ไร่)','% ของ Pure A403'],[[r['soil_unit'],names.get(r['soil_unit'],'ไม่พบชื่อในตารางอ้างอิง'),f"{r['area_rai']:,.4f}",f"{r['share_of_all_pure_pct']:.6f}"] for r in top]+[['Uncovered','พื้นที่ไม่ตัดกับดิน',f'{unc:.6e}',f'{100*unc/den:.6e}']])
    annualtable=table(['ปี พ.ศ. / ค.ศ.','Yield กก./ไร่','T2M °C','ฝน มม.','Production ตัน','Bearing ไร่'],[[f'{int(r["year_ce"])+543} / {r["year_ce"]}',f'{float(r["yield_kg_per_rai"]):,.0f}',f'{float(r["weather_t2m_mean_c"]):.2f}',f'{float(r["weather_precip_total_mm"]):,.1f}',f'{float(r["production_tonnes"]):,.0f}',f'{float(r["bearing_area_rai"]):,.0f}'] for r in annual])
    sources_used=[P9/'province_year.csv',P9/'weather_profile_month.csv',P9/'sampling_points.csv',P9/'source_manifest.json',P9/'field_lineage.jsonl',P9/'output_manifest.json',P10/'dataset.json',P10/'qc_summary.json',P10/'release_metadata.json',P10/'output_manifest.json',P11/'district_year.json',P11/'source_manifest.json',P11/'field_lineage.json',P11/'output_manifest.json',workbook]
    inputhash={str(p.relative_to(ROOT)):sha(p) for p in sources_used}
    source_list=''.join('<li><code>'+e(k)+'</code><br><small>SHA256 '+v+'</small></li>' for k,v in inputhash.items())
    html=f'''<!doctype html><html lang="th"><meta charset="utf-8"><title>EDA จันทบุรีสำหรับหารืออาจารย์</title><style>
    *{{box-sizing:border-box}}body{{margin:0;background:#f4f7f7;color:#132d38;font:18px/1.75 "Leelawadee UI",Tahoma,sans-serif}}main{{max-width:1100px;margin:auto;padding:36px 42px;background:white}}h1{{font-size:34px;line-height:1.4}}h2{{margin-top:46px;font-size:26px}}h3{{font-size:20px}}p{{margin:13px 0}}small,.note{{font-size:15px}}svg{{display:block;width:100%;height:auto;margin:15px 0 28px;font:16px "Leelawadee UI",Tahoma,sans-serif}}svg text{{fill:#132d38}}circle:focus{{stroke:#000;stroke-width:3}}table{{border-collapse:collapse;width:100%;font-size:15px}}td,th{{padding:10px;border-bottom:1px solid #dbe3e9;text-align:left}}td:nth-last-child(-n+3){{text-align:right}}th{{background:#eef3f6}}.table-scroll{{overflow-x:auto}}code{{overflow-wrap:anywhere;font-size:13px}}li{{margin:12px 0}}details{{margin:18px 0}}summary{{cursor:pointer;font-weight:bold}}button{{font:inherit;padding:8px 16px;background:white;border:1px solid #889aa4}}@media(max-width:650px){{main{{padding:20px}}h1{{font-size:27px}}}}@media print{{body{{background:white}}main{{padding:0;max-width:none}}button{{display:none}}svg,table,tr{{break-inside:avoid}}h2,h3{{break-after:avoid}}details>div{{display:block}}}}@page{{size:A4;margin:15mm}}
    </style><main><p class="note">Internal Research Only · EDA สำหรับคุยกับอาจารย์ · 12 ตุลาคม 2569 · Public HOLD</p><h1>จันทบุรี ผลผลิต อากาศ และความหลากหลายของดิน</h1>
    <p>อ่านข้อมูลเดิมแบบ Read-only ไม่ฝึกโมเดล ไม่ทำนาย ไม่ตีพิมพ์ และไม่สรุปเหตุและผล หน่วยเวลาของกราฟจังหวัดคือหนึ่งจังหวัด–ปี รวม20ปี ค.ศ.2006–2025 (พ.ศ.2549–2568) ไม่ใช่20สวนหรือ20ตัวอย่างอิสระอย่างแน่นอน</p><button onclick="window.print()">พิมพ์ / บันทึก PDF</button>
    <h2>1 แนวโน้มจังหวัดรายปี</h2><p>อากาศใช้ <strong>มกราคม–ธันวาคมของปีเดียวกัน</strong> อุณหภูมิเป็นค่าเฉลี่ยถ่วงจำนวนวัน ฝนเป็นผลรวมทั้งปี ไม่ใช่เฉพาะออกดอกหรือติดผล กราฟแยกแกนเพื่อไม่ให้ความคล้ายของเส้นจากแกนคู่ชวนตีความเกินข้อมูล</p>{trend}
    <details><summary>ดูตารางจังหวัด–ปีทั้ง20ปี</summary>{annualtable}</details>
    <h2>2 อุณหภูมิและฝนสัมพันธ์กับ Yield อย่างไรในข้อมูลนี้</h2><p>หนึ่งจุดคือหนึ่งจังหวัด–ปี ใช้ช่วงเดือนและข้อมูลชุดเดียวกับกราฟแนวโน้ม ตัวเลขกำกับเลือกปี2016/2022/2024เพื่ออ่านกรณีตัวอย่าง ไม่ใช่ตัดปีอื่นออก Hover/Focusจุดเพื่อดูปีและค่า ทุกค่าดูได้ในตารางด้านบน ไม่มีเส้นโมเดลพยากรณ์</p>{scatter}<p>Pearson r: อุณหภูมิ {rt:.3f}; ฝน {rr:.3f}; n=20 เป็นความสัมพันธ์เชิงพรรณนาของข้อมูลดิบ ไม่รายงานp-valueหรือช่วงความเชื่อมั่น ไม่ควบคุมTrend/Serial Dependence และConfounders จึงไม่ใช้เป็นหลักฐานว่าร้อนหรือฝนทำให้ผลผลิตเปลี่ยน</p>
    <h2>3 Yield สูงและต่ำในสิบอำเภอ</h2><p>เรียงYieldจากสูงไปต่ำแยกปี พ.ศ.2566/2567 หน่วยกก./ไร่ของพื้นที่ให้ผล ไม่ใช่Productionรวม <strong>รับรองเฉพาะถอดตาราง Phase11B ยังไม่รับรองFinal Actual นิยามปีและPublicationVintage</strong> ผลผลิตปริมาณรวมกับBearingแสดงไว้เพื่อไม่สับสนระหว่างพื้นที่มากกับผลิตภาพสูง</p>{table(['ปี พ.ศ.','อันดับภายในปี','อำเภอ','Yield กก./ไร่','Production ตัน','Bearing ไร่'],rank_data)}
    <h2>4 ดินใน Pure A403 อำเภอท่าใหม่เท่านั้น</h2><p>ตัวหาร {den:,.2f} ไร่ คือPure A403ที่ตัดขอบเขตท่าใหม่แล้ว ไม่รวมMixed-useที่ไม่ทราบสัดส่วนทุเรียน มี43รหัสหน่วยแผนที่ SCและWอยู่ในCoveredอยู่แล้ว ไม่บวกซ้ำ UncoveredมีเพียงNumerical Residual {unc:.3e} ไร่ ยังเก็บค่าเดิม</p>{''.join(bar)}
    <p class="note">Otherใช้จัดกลุ่มแสดงผลเท่านั้น ไม่แก้43รหัสต้นฉบับ SCเป็นพื้นที่ลาดชันเชิงซ้อน Wเป็นน้ำ ไม่ใช่ชื่อชุดดินเดี่ยว ทศนิยม0.000%ในภาพไม่แปลว่าเป็นศูนย์จริง ให้ดูตารางละเอียด</p><details><summary>ดูทุกหน่วยแผนที่ดินและพื้นที่เดิม</summary>{soiltable}</details>
    <p>Land-use Map พ.ศ.2568; Soil Production Year พ.ศ.2561; CRS EPSG:32647 พื้นที่ไร่=m²/1,600 Polygonคือหน่วยการใช้ที่ดิน ไม่ใช่สวนรายเจ้าของ วันสำรวจเฉพาะพื้นที่ยังไม่ยืนยัน และแผนที่2568ไม่พิสูจน์ว่าพื้นที่ปลูกเหมือนเดิมย้อนหลังทุกปี ไม่เปรียบเทียบSoil–Yieldข้ามสิบอำเภอจากดินท่าใหม่อำเภอเดียว</p><p class="note">Phase10F Internal Accepted Proposal ภายใต้QC Policy v2, toleranceเดิม0.1ไร่; GEOS3.13.1มีTwo-stage Union Warningเติมรูประมาณ14.2778ไร่ในDiagnostic Path แต่Primary Policyผ่าน ใช้พื้นที่/สัดส่วนจากReleaseเดิมโดยไม่UnionหรือOverlayใหม่ ไม่อ้างCross-engine Certification และPublic HOLDคงเดิม</p>
    <h2>5 ข้อสังเกตสำหรับเปิดบทสนทนา</h2><ol>{''.join('<li>'+e(f)+'</li>' for f in findings)}</ol>
    <h2>คำถามที่ควรนำไปถามอาจารย์</h2><ol><li>ควรใช้ช่วงอากาศทั้งปีเป็นภาพรวมก่อน แล้วกำหนดช่วงออกดอก–ติดผลจากหลักฐานใด? ฤดูอาจข้ามปีและต่างการจัดการ ห้ามเลือกเดือนย้อนหลังเพื่อไล่Correlation</li><li>Yieldปี2567ที่ลดในทุกอำเภอสะท้อนสภาพฤดูหรือความต่างSourceVintage/นิยามอย่างไร? ควรขอข้อมูลพันธุ์ อายุต้น การให้น้ำ และความเสียหายเพิ่มหรือไม่?</li><li>หากจะศึกษาอำเภอ ควรหาYieldหลายปีและDistrict Weather Exposureที่ตรวจได้ก่อนขยายดินอีก9อำเภอหรือไม่?</li><li>จะควบคุมแนวโน้มเวลา พื้นที่ให้ผล และSpatial/Temporal Dependenceอย่างไร ก่อนพิจารณาว่าข้อมูลพอวิเคราะห์หรือยัง?</li><li>ยอมรับStatic Footprint2568เป็นRetrospective Proxyเพื่อการสำรวจหรือจำเป็นต้องมีประวัติพื้นที่ปลูก และต้องขอสิทธิ์ใช้Derived GISใดก่อนเผยแพร่?</li></ol>
    <h2>แหล่งข้อมูล วิธีคำนวณ และข้อจำกัด</h2><p>OAEจังหวัดจากWorkbookเดิมตรวจค่าและCell Anchorsครบ80ค่ากับPhase9D ใช้Yieldที่ต้นทางรายงาน ไม่คำนวณเติม ไม่แจกYieldจังหวัดลงสวน/Polygon; รายอำเภอจากPDFเดิมหน้า5Table2ของPhase11B; ดินจากPhase10F ไม่ทำOverlayใหม่</p>
    <p>POWERเป็นอากาศแบบGrid/Spatial Profile ไม่ใช่เซนเซอร์สวนหรือสถานีอิสระ เอกสารเดิมระบุMeteorologyประมาณ0.5°×0.625° (ระดับหลายสิบกิโลเมตร) SolarคนละResolution; กรอบSampling0.25°เป็นการแบ่งชั้นของโครงการ ไม่ใช่Native Grid จันทบุรี16จุดมีเพียง6Meteorology Response Signatures และ11Combined Weather/Solar Profiles Native Grid IDsยังไม่ยืนยัน จึงไม่ตีความความต่างทุกอำเภอหรือสวนเล็กอย่างอิสระ</p>
    <p>Weather Profileถ่วงพื้นที่A403ตามFootprint2568และWeightเดิมรวม1 อุณหภูมิรายปี=Σ(T2Mเดือน×จำนวนวัน)/Σวัน ฝนรายปี=Σ(PRECTOTCORRมม./วัน×วัน) ตรวจครบ12เดือนและCoverageเดิม ไม่เติมMissing Profileเดือนเป็นปีเฉลี่ยแทน NASAHeaderเดิมใช้LST; OAEนิยามปี/PublicationVintageและRevisionsยังมีข้อจำกัด จับคู่ProvinceCode22/ปีที่Sourceระบุ ไม่รับรองSeason AlignmentหรือHistorical PIT อากาศครบทั้งปีเป็นข้อมูลย้อนหลัง ไม่ใช่สิ่งที่ทราบก่อนฤดู</p>
    <p>สัดส่วนดิน=Areaของแต่ละหน่วยที่ตัดPure A403ท่าใหม่÷Pure Areaทั้งอำเภอ×100 รวมSC/W/Uncoveredในตัวหาร แต่Excluded Mixedไม่อยู่ในตัวหาร Pearson r=Σ[(x−mean x)(y−mean y)]/√[Σ(x−mean x)²Σ(y−mean y)²] ไม่แปลCorrelationเป็นCausal Impact</p>
    <details><summary>Source Paths และ SHA256 สำหรับตรวจย้อนกลับ</summary><ul>{source_list}</ul></details><p class="note">บันทึกวิธี/ค่าที่ใช้ใน analysis_summary.json และ source_manifest.json ภายในโฟลเดอร์รายงาน ตรวจSeal/Protected Hashก่อน–หลัง คงPublic HOLDเพราะสิทธิ์Derived LDDยังไม่ยืนยัน ไม่มีNetwork/Training/APIChanges</p>
    </main></html>'''
    assert all(sha(ROOT/k)==v for k,v in before.items()),'Protected changed'
    OUT.mkdir(parents=True)
    with (OUT/'report_th.html').open('x',encoding='utf-8') as f:f.write(html)
    save(OUT/'analysis_summary.json',dict(coverage_year_ce=yy,weather_months=list(range(1,13)),
        pearson_temperature=rt,pearson_rain=rr,findings=findings,district_ranking=rank_data,
        soil_denominator_rai=den,uncovered_rai=unc,top3_soil_share_pct=top3,soil=soil,
        oae_workbook_cell_anchors=source,weather_qc='240 months recomputed; all matched annual within1e-6',
        interpretation='descriptive only, no p-values, no causal or predictive claim'))
    save(OUT/'source_manifest.json',dict(input_hashes=inputhash,code_sha256=sha(Path(__file__)),
        protected_files_checked=len(before),protected_hash_mismatches=[],production_sha256=sha(MODEL),
        sources_read_only=True,public_release='HOLD',historical_pit='not_certified'))
    files={p.name:sha(p) for p in OUT.iterdir() if p.is_file()}
    save(OUT/'output_manifest.json',dict(files=files,manifest_payload_sha256=hashlib.sha256(json.dumps(files,sort_keys=True,ensure_ascii=False).encode()).hexdigest(),public_release='HOLD'))
    print(json.dumps({'report':str(OUT/'report_th.html'),'findings':findings,'protected':len(before)},ensure_ascii=False))


if __name__=='__main__':
    assert corr([1,2,3],[2,4,6])==1 and corr([1,2,3],[6,4,2])==-1
    build()
